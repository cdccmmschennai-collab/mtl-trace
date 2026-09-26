"""Shared tabular extraction (TRD §12.1) — one implementation for every source type.

All supplied source documents are worker-export workbooks: a join-key column holding the milestone tag,
"value as per document" columns, and document/page provenance columns. Which of those a document has,
and under which header wording, varies. So an adapter is only *identification* (which sheet belongs to
its document type); field discovery is shared and keyword-based (`backend.domain.fields`):

    header cell → classify_header → join key / attribute / provenance / not extracted
        join key   exactly one column → bound · none → MISSING_FIELD (error) · several → AMBIGUOUS_FIELD (error)
        attribute  exactly one column → bound · none → FIELD_NOT_PROVIDED (warning) · several → AMBIGUOUS_FIELD (error)
        provenance every matching column is kept

Available-data principle: a document carries the attributes it actually has. An attribute column the
document does not have is reported as not provided — never filled, never guessed — and the document still
processes. A document with no attribute column at all has nothing to contribute (NO_ATTRIBUTE_FIELD).

Nothing is ever guessed. Column letters appear only in the *reported* mapping and in provenance.
"""

import io
import zipfile
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils import get_column_letter

from backend.domain.fields import (
    ACCEPTED_TAG_KEY,
    ATTRIBUTE_LABELS,
    JOIN_KEY,
    PROVENANCE_ROLES,
    FieldKind,
    HeaderMeaning,
    classify_header,
    contains_phrase,
)
from backend.domain.headers import normalize_header
from backend.domain.issues import Issue, Severity
from backend.domain.sources import ATTRIBUTES, Attribute
from backend.domain.tags import normalize_tag
from backend.domain.values import cell_text

HEADER_SEARCH_ROWS = 20
SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm")


# ------------------------------------------------------------------ configuration

@dataclass(frozen=True)
class TabularAdapterConfig:
    """Identification of one document type. Evidence, in order of strength:
    - `sheet_names` / `evidence_headers`: exact texts verified on a real file of the type;
    - `keywords`: whole-word phrases naming the type in a sheet title or in the header row;
    - the user's explicit choice of type, for a type whose real format is not yet verified
      (`verified=False`): the one sheet with a TAG NUMBER header row is used."""
    source_type: str
    document_label: str                       # "Data Sheet"
    sheet_names: tuple[str, ...] = ()
    evidence_headers: tuple[str, ...] = ()
    keywords: tuple[str, ...] = ()
    verified: bool = True


# ------------------------------------------------------------------ results

@dataclass(frozen=True)
class BoundField:
    key: str                                  # TAG_KEY · TAG_AS_PER_DOC · MAKE … · DOCUMENT_REFERENCE · PAGE_NO
    label: str
    column: str                               # letter, for reporting and provenance only
    header: str                               # header text exactly as found


@dataclass(frozen=True)
class ExtractedValue:
    raw: str | None                           # exact cell content as text (surrounding spaces kept)
    normalized: str | None                    # trimmed; whitespace-only → None
    column: str
    raw_type: str | None


@dataclass(frozen=True)
class ExtractedRow:
    row_number: int
    raw_key: str | None
    normalized_key: str | None
    values: dict[Attribute, ExtractedValue]   # only the attributes the document provides
    provenance: dict[str, str]                # populated provenance cells, by header as found (stripped)


@dataclass
class Identification:
    source_type: str | None = None
    sheet_name: str | None = None
    skipped_sheets: list[dict] = field(default_factory=list)     # [{sheet, reason}]
    issues: list[Issue] = field(default_factory=list)


@dataclass
class Extraction:
    source_type: str
    sheet_name: str
    header_row: int | None = None
    mapping: list[BoundField] = field(default_factory=list)
    headers_found: list[str] = field(default_factory=list)
    rows: list[ExtractedRow] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    provided_attributes: tuple[Attribute, ...] = ()               # bound attributes, in Attribute order
    provenance_roles: dict[str, str] = field(default_factory=dict)  # provenance header -> role

    @property
    def ok(self) -> bool:
        return not any(i.severity == Severity.ERROR for i in self.issues)

    def location_of(self, row: ExtractedRow) -> dict:
        """Provenance of one row: each role (document_reference / document_idb / page), joining several
        columns of one role in column order so no reference is dropped or chosen, plus every populated
        provenance cell under its own header."""
        loc: dict = {}
        for role in PROVENANCE_ROLES:
            vals = [row.provenance[h] for h, r in self.provenance_roles.items() if r == role and h in row.provenance]
            loc[role] = "; ".join(vals) if vals else None
        loc["source_fields"] = dict(row.provenance)
        return loc


# ------------------------------------------------------------------ workbook loading

@dataclass
class Workbook:
    file_name: str
    sheets: dict[str, list[tuple]]            # sheet title -> rows of cell values


class UnreadableWorkbook(Exception):
    pass


def load_workbook(file_name: str, data: bytes) -> Workbook:
    if not file_name.lower().endswith(SUPPORTED_EXTENSIONS):
        raise UnreadableWorkbook(
            f"'{file_name}' is not an Excel workbook (.xlsx / .xlsm). Legacy .xls files must be saved as .xlsx first.")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise UnreadableWorkbook(f"'{file_name}' could not be opened as a workbook: {exc}") from exc
    try:
        return Workbook(file_name, {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets})
    finally:
        wb.close()


# ------------------------------------------------------------------ adapter

def _err(code: str, message: str, **kw) -> Issue:
    return Issue(Severity.ERROR, code, message, **kw)


def _warn(code: str, message: str, **kw) -> Issue:
    return Issue(Severity.WARNING, code, message, **kw)


def _raw_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value if value != "" else None
    return cell_text(value)


def _is_join_key(value: object) -> bool:
    m = classify_header(value)
    return m is not None and m.kind == FieldKind.JOIN_KEY


def header_row(rows: list[tuple]) -> int | None:
    """Index of the first row (within the search window) holding a TAG NUMBER join-key header."""
    return next((i for i, row in enumerate(rows[:HEADER_SEARCH_ROWS]) if any(_is_join_key(v) for v in row)), None)


class TabularAdapter:
    def __init__(self, config: TabularAdapterConfig):
        self.config = config
        self._sheet_names = {normalize_header(n) for n in config.sheet_names}
        self._evidence = {normalize_header(h) for h in config.evidence_headers}
        self._keywords = tuple(normalize_header(k) for k in config.keywords)

    @property
    def source_type(self) -> str:
        return self.config.source_type

    # ---- identification

    def _named(self, text: object) -> bool:
        t = normalize_header(text)
        return any(contains_phrase(t, k) for k in self._keywords)

    def evidence_sheets(self, book: Workbook) -> list[str]:
        """Sheets carrying this document type's identification evidence (never data cells for keywords)."""
        hits = []
        for name, rows in book.sheets.items():
            if normalize_header(name) in self._sheet_names or self._named(name):
                hits.append(name)
                continue
            if any(normalize_header(v) in self._evidence for row in rows[:HEADER_SEARCH_ROWS] for v in row):
                hits.append(name)
                continue
            header = header_row(rows)
            if header is not None and any(self._named(v) for v in rows[header]):
                hits.append(name)
        return hits

    def identify(self, book: Workbook, requested: bool = False) -> Identification:
        hits = self.evidence_sheets(book)
        ident = Identification()
        label = self.config.document_label
        if not hits and requested and not self.config.verified:
            # No real sample of this type has been verified: the user's explicit choice of type is the evidence,
            # provided exactly one sheet has a TAG NUMBER header row to extract.
            hits = [n for n, rows in book.sheets.items() if header_row(rows) is not None]
        if not hits:
            named = " / ".join([*self.config.sheet_names, *self.config.keywords])
            ident.issues.append(_err(
                "NOT_THIS_SOURCE_TYPE",
                f"'{book.file_name}' does not look like a {label}: no sheet title or header naming it ({named})."))
        elif len(hits) > 1:
            ident.issues.append(_err(
                "SHEET_AMBIGUOUS",
                f"'{book.file_name}': several sheets look like {label} extracts "
                f"({', '.join(repr(h) for h in hits)}); the sheet to extract cannot be chosen safely."))
        else:
            ident.source_type = self.source_type
            ident.sheet_name = hits[0]
        ident.skipped_sheets = [
            {"sheet": n, "reason": f"no {label} evidence" if n not in hits else "ambiguous"}
            for n in book.sheets if n != ident.sheet_name]
        return ident

    # ---- extraction

    def extract(self, book: Workbook, sheet_name: str) -> Extraction:
        return extract_sheet(self.source_type, book, sheet_name)


def extract_sheet(source_type: str, book: Workbook, sheet_name: str) -> Extraction:
    """Discover fields by header meaning and extract every row. Shared by all source types."""
    rows = book.sheets[sheet_name]
    ext = Extraction(source_type, sheet_name)
    doc = f"'{book.file_name}', sheet '{sheet_name}'"

    header_idx = header_row(rows)
    if header_idx is None:
        found = next(([str(v).strip() for v in r if cell_text(v)] for r in rows[:HEADER_SEARCH_ROWS]
                      if any(cell_text(v) for v in r)), [])
        ext.headers_found = found
        ext.issues.append(_err(
            "MISSING_FIELD",
            f"{doc}: no header row contains the join key TAG NUMBER (accepted: {ACCEPTED_TAG_KEY}). "
            f"Headers found: {', '.join(found) or '(none)'}.", value="TAG NUMBER (join key)"))
        return ext

    header = rows[header_idx]
    ext.header_row = header_idx + 1
    ext.headers_found = [str(v) for v in header if cell_text(v) is not None]
    headers_list = ", ".join(repr(h) for h in ext.headers_found)

    # Classify every header once; group columns by the concept they carry.
    by_key: dict[str, list[tuple[int, HeaderMeaning]]] = {}
    for i, v in enumerate(header):
        m = classify_header(v)
        if m is not None:
            by_key.setdefault(m.key, []).append((i, m))

    def bind(i: int, key: str, label: str) -> None:
        ext.mapping.append(BoundField(key, label, get_column_letter(i + 1), str(header[i])))

    def ambiguous(label: str, cands: list[tuple[int, HeaderMeaning]]) -> None:
        where = ", ".join(f"{get_column_letter(i + 1)} {str(header[i])!r}" for i, _ in cands)
        ext.issues.append(_err(
            "AMBIGUOUS_FIELD",
            f"{doc}: {label} matches {len(cands)} columns ({where}); refusing to choose one. "
            f"Headers found: {headers_list}.", value=label))

    join_cands = by_key.get(JOIN_KEY, [])
    if len(join_cands) > 1:
        ambiguous("TAG NUMBER (join key)", join_cands)
    else:
        bind(join_cands[0][0], JOIN_KEY, "TAG NUMBER (join key)")
    join_col = join_cands[0][0]

    attr_cols: dict[Attribute, int] = {}
    not_provided: list[str] = []
    for attribute in ATTRIBUTES:
        label = ATTRIBUTE_LABELS[attribute]
        cands = [(i, m) for i, m in (c for cs in by_key.values() for c in cs) if m.attribute == attribute]
        if len(cands) == 1:
            attr_cols[attribute] = cands[0][0]
            bind(cands[0][0], cands[0][1].key, label)
        elif not cands:
            not_provided.append(label)
        else:
            ambiguous(label, cands)

    prov_cols: list[int] = []
    for key in ("DOCUMENT_REFERENCE", "DOCUMENT_IDB", "PAGE_NO"):
        for i, m in by_key.get(key, []):
            bind(i, key, str(header[i]).strip())
            prov_cols.append(i)
    prov_cols.sort()
    ext.provenance_roles = {str(header[i]).strip(): classify_header(header[i]).role for i in prov_cols}

    if ext.ok and not attr_cols:
        ext.issues.append(_err(
            "NO_ATTRIBUTE_FIELD",
            f"{doc}: no MTL attribute column found (TAG NUMBER as per document, MAKE, MODEL, SERIAL NUMBER, "
            f"PART NUMBER), so the document has nothing to contribute. Headers found: {headers_list}."))
    if not ext.ok:
        return ext
    if not_provided:
        ext.issues.append(_warn(
            "FIELD_NOT_PROVIDED",
            f"{doc}: no column for {', '.join(not_provided)}. This document does not provide "
            f"{'that attribute' if len(not_provided) == 1 else 'those attributes'}; nothing is filled in. "
            f"Headers found: {headers_list}.", value=", ".join(not_provided)))
    ext.provided_attributes = tuple(a for a in ATTRIBUTES if a in attr_cols)

    body = rows[header_idx + 1:]
    while body and all(cell_text(v) is None for v in body[-1]):
        body.pop()

    def cell(row: tuple, idx: int):
        return row[idx] if idx < len(row) else None

    empty_rows: list[int] = []
    non_text_keys: list[int] = []
    for offset, row in enumerate(body):
        row_number = ext.header_row + 1 + offset
        if all(cell_text(v) is None for v in row):
            empty_rows.append(row_number)
            continue
        key_value = cell(row, join_col)
        key_text = cell_text(key_value)
        if key_text is not None and not isinstance(key_value, str):
            non_text_keys.append(row_number)
        values = {}
        for attribute, idx in attr_cols.items():
            v = cell(row, idx)
            values[attribute] = ExtractedValue(raw=_raw_text(v), normalized=cell_text(v),
                                               column=get_column_letter(idx + 1),
                                               raw_type=type(v).__name__ if v is not None else None)
        provenance = {str(header[i]).strip(): t for i in prov_cols if (t := cell_text(cell(row, i))) is not None}
        ext.rows.append(ExtractedRow(
            row_number=row_number, raw_key=_raw_text(key_value),
            normalized_key=normalize_tag(key_text) if key_text is not None else None,
            values=values, provenance=provenance))

    if empty_rows:
        ext.issues.append(_warn("EMPTY_ROW", f"{len(empty_rows)} completely empty row(s) were skipped.",
                                rows=tuple(empty_rows)))
    if non_text_keys:
        ext.issues.append(_warn(
            "NON_TEXT_TAG_KEY",
            f"{len(non_text_keys)} TAG NUMBER (join key) value(s) are numbers or dates rather than text.",
            rows=tuple(non_text_keys)))
    return ext
