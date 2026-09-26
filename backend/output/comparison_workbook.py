"""Comparison workbook — generated, never cloned (TRD §20).

Layout follows the mental-model reference `5.MTL-DATA MISMATCH COMPARISION - AUTOMATED.xlsx`
(structure only — none of its data) and the working file's `SUMMERY` shell:

    SUMMERY
    <ATTRIBUTE> COMPARISION  × one per attribute
        A S.NO · B TAG NUMBER · C TAG DISCRIPTION · D EQUIPMENT DESCRIPTION · E SIZE & RATING
        one column per PARTICIPATING source type, "{attribute label} IN {source label}", display order
        REMARKS · FINAL <attribute> · STATUS

Sheets and source columns come from the model, so a new source type (GA) or attribute extends the
workbook with no template and no code change. A source type not in the run gets no column anywhere.

The consolidated workbook (Phase 1C) carries source values only; REMARKS, FINAL, STATUS and the SUMMERY
result counts are blank. The comparison workbook (Phase 1D) is the same layout filled from the comparison
result. The workbook is re-opened and validated before it is handed back; a failed guard raises.
"""

import io
from dataclasses import dataclass, field

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SCOPE_HEADERS = ("S.NO", "TAG NUMBER", "TAG DISCRIPTION", "EQUIPMENT DESCRIPTION", "SIZE & RATING")
SUMMARY_SHEET = "SUMMERY"
SUMMARY_LEADING = ("SL NO", "DESCRIPTION")
# The three comparison states (the working file's legacy "PENDING/NOT AVAILABLE" header is not a state and
# is not reproduced; rows without any source value have no verdict and are simply not counted).
SUMMARY_RESULT_HEADERS = ("MATCH", "MISMATCH", "REVIEW REQUIRED")
TOTAL_ROW_LABEL = "TOTAL EQUIPMENT / TAG COUNT"
ALLOWED_STATUS = {"MATCH", "MISMATCH", "REVIEW REQUIRED"}

PURPLE = PatternFill("solid", fgColor="FF7030A0")
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(name="Calibri", size=11, bold=True)
SCOPE_WIDTHS = (5.6, 26.3, 45.1, 44.9, 43.7)


# ---------------------------------------------------------------- model

@dataclass(frozen=True)
class SourceColumn:
    source_type: str
    label: str                          # business label, used in "{attribute} IN {label}"
    summary_code: str                   # SUMMERY "<summary_code> QTY" header


@dataclass(frozen=True)
class ScopeRow:
    s_no: str | None
    tag_number: str
    tag_description: str | None
    equipment_description: str | None
    size_rating: str | None


@dataclass(frozen=True)
class SourceCell:
    value: str | None = None            # canonical raw value; None = nothing to write
    note: str | None = None             # e.g. duplicate candidates — written as a cell comment, not a value


@dataclass
class AttributeSheet:
    attribute: str
    sheet_name: str
    column_label: str
    final_label: str
    summary_description: str
    cells: list[dict[str, SourceCell]]  # one dict per scope row: source_type -> cell


@dataclass
class WorkbookModel:
    scope: list[ScopeRow]
    sources: list[SourceColumn]
    sheets: list[AttributeSheet]
    title: str = ""
    # Per attribute, one entry per scope row. Empty dicts (consolidated workbook) leave every row blank.
    remarks: dict[str, list[str | None]] = field(default_factory=dict)
    finals: dict[str, list[str | None]] = field(default_factory=dict)
    statuses: dict[str, list[str | None]] = field(default_factory=dict)

    def column(self, which: dict[str, list[str | None]], sheet: "AttributeSheet") -> list[str | None]:
        return which.get(sheet.attribute) or [None] * len(self.scope)

    def result_counts(self, sheet: "AttributeSheet") -> list[int | None]:
        """SUMMERY result counts; blank when the sheet has no comparison result at all."""
        if sheet.attribute not in self.statuses:
            return [None] * len(SUMMARY_RESULT_HEADERS)
        statuses = self.statuses[sheet.attribute]
        return [sum(1 for s in statuses if s == h) for h in SUMMARY_RESULT_HEADERS]

    def source_header(self, sheet: AttributeSheet, source: SourceColumn) -> str:
        return f"{sheet.column_label} IN {source.label}"

    def sheet_headers(self, sheet: AttributeSheet) -> list[str]:
        return [*SCOPE_HEADERS, *(self.source_header(sheet, s) for s in self.sources),
                "REMARKS", sheet.final_label, "STATUS"]

    def summary_headers(self) -> list[str]:
        return [*SUMMARY_LEADING, *(f"{s.summary_code} QTY" for s in self.sources), "REMARKS", *SUMMARY_RESULT_HEADERS]

    def qty(self, sheet: AttributeSheet, source_type: str) -> int:
        """Per-source non-empty count, computed from the same values written into the sheet."""
        return sum(1 for row in sheet.cells if row[source_type].value not in (None, ""))


class WorkbookValidationError(Exception):
    def __init__(self, problems: list[str]):
        super().__init__("Generated workbook failed validation: " + "; ".join(problems[:10]))
        self.problems = problems


# ---------------------------------------------------------------- writer

def build_workbook(model: WorkbookModel) -> bytes:
    _check_model(model)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if model.title:
        wb.properties.title = model.title
    _write_summary(wb.create_sheet(SUMMARY_SHEET), model)
    for sheet in model.sheets:
        _write_attribute_sheet(wb.create_sheet(sheet.sheet_name), model, sheet)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    problems = validate_workbook(data, model)
    if problems:
        raise WorkbookValidationError(problems)
    return data


def _check_model(model: WorkbookModel) -> None:
    codes = [s.source_type for s in model.sources]
    if len(set(codes)) != len(codes):
        raise ValueError("A source type appears twice in the workbook model.")
    for sheet in model.sheets:
        if len(sheet.cells) != len(model.scope):
            raise ValueError(f"{sheet.sheet_name}: {len(sheet.cells)} rows for {len(model.scope)} scope tags.")
        for row in sheet.cells:
            if set(row) != set(codes):
                raise ValueError(f"{sheet.sheet_name}: row source set {sorted(row)} differs from the run's {sorted(codes)}.")
        for which in (model.remarks, model.finals, model.statuses):
            if sheet.attribute in which and len(which[sheet.attribute]) != len(model.scope):
                raise ValueError(f"{sheet.sheet_name}: comparison column length differs from the scope.")
        bad = {s for s in model.statuses.get(sheet.attribute, []) if s is not None} - ALLOWED_STATUS
        if bad:
            raise ValueError(f"{sheet.sheet_name}: {sorted(bad)} are not comparison states.")


def _header(ws, row: int, col: int, value: str, fill: PatternFill, center: bool = True) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.fill = fill
    c.font = BOLD
    c.border = BORDER
    c.alignment = Alignment(horizontal="center" if center else None, vertical="bottom")


def _data(ws, row: int, col: int, value) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.border = BORDER
    if isinstance(value, str):
        c.data_type = "s"                   # always literal text: a source value starting with '=' is not a formula


def _write_attribute_sheet(ws, model: WorkbookModel, sheet: AttributeSheet) -> None:
    headers = model.sheet_headers(sheet)
    n_scope, n_src = len(SCOPE_HEADERS), len(model.sources)
    for col, h in enumerate(headers, start=1):
        _header(ws, 1, col, h, PURPLE if col <= n_scope else YELLOW, center=col <= n_scope + n_src)
    remarks = model.column(model.remarks, sheet)
    finals = model.column(model.finals, sheet)
    statuses = model.column(model.statuses, sheet)
    for i, (scope, cells) in enumerate(zip(model.scope, sheet.cells)):
        r = i + 2
        for col, v in enumerate((scope.s_no, scope.tag_number, scope.tag_description,
                                 scope.equipment_description, scope.size_rating), start=1):
            _data(ws, r, col, v)
        for j, src in enumerate(model.sources):
            cell = cells[src.source_type]
            _data(ws, r, n_scope + 1 + j, cell.value)
            if cell.note:
                ws.cell(row=r, column=n_scope + 1 + j).comment = Comment(cell.note, "CDC MTL Tool")
        base = n_scope + n_src
        _data(ws, r, base + 1, remarks[i])
        _data(ws, r, base + 2, finals[i])
        _data(ws, r, base + 3, statuses[i])
    for col, w in enumerate(SCOPE_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    for j, h in enumerate(headers[n_scope:], start=n_scope + 1):
        ws.column_dimensions[get_column_letter(j)].width = max(18.0, min(len(h) + 4.0, 45.0))
    ws.column_dimensions[get_column_letter(n_scope + n_src + 1)].width = 40.0
    ws.freeze_panes = "A2"


def _write_summary(ws, model: WorkbookModel) -> None:
    headers = model.summary_headers()
    for col, h in enumerate(headers, start=1):
        _header(ws, 1, col, h, YELLOW, center=False)
        ws.cell(row=1, column=col).font = Font(name="Calibri", size=11, bold=False)
    for i, sheet in enumerate(model.sheets, start=1):
        r = i + 1
        _data(ws, r, 1, i)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        _data(ws, r, 2, sheet.summary_description)
        for j, src in enumerate(model.sources):
            _data(ws, r, 3 + j, model.qty(sheet, src.source_type))     # static value, not a formula (Q7)
        base = 3 + len(model.sources)
        _data(ws, r, base, None)                                       # REMARKS
        for k, count in enumerate(model.result_counts(sheet), start=1):
            _data(ws, r, base + k, count)
    total_row = len(model.sheets) + 3
    ws.cell(row=total_row, column=2, value=TOTAL_ROW_LABEL)
    ws.cell(row=total_row, column=3, value=len(model.scope))
    ws.column_dimensions["A"].width = 6.1
    ws.column_dimensions["B"].width = 30.0
    for col in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = max(9.5, len(headers[col - 1]) + 2.0)


# ---------------------------------------------------------------- validation (mandatory before saving)

def validate_workbook(data: bytes, model: WorkbookModel) -> list[str]:
    """Re-open the generated bytes and check them against the model. Returns every problem found."""
    problems: list[str] = []
    wb = openpyxl.load_workbook(io.BytesIO(data))
    expected_sheets = [SUMMARY_SHEET, *(s.sheet_name for s in model.sheets)]
    if wb.sheetnames != expected_sheets:
        return [f"sheets are {wb.sheetnames}, expected {expected_sheets}"]
    n = len(model.scope)
    n_scope, n_src = len(SCOPE_HEADERS), len(model.sources)
    participating = {s.label for s in model.sources}

    for sheet in model.sheets:
        ws = wb[sheet.sheet_name]
        headers = [c.value for c in ws[1]]
        if headers != model.sheet_headers(sheet):
            problems.append(f"{sheet.sheet_name}: headers {headers} differ from the expected layout")
            continue
        # Guard 2 — no column for a source type that did not participate.
        for h in headers[n_scope:n_scope + n_src]:
            if h.split(" IN ", 1)[1] not in participating:
                problems.append(f"{sheet.sheet_name}: column {h!r} is not a participating source")
        # Guard 3 — row count and scope order.
        if ws.max_row != n + 1:
            problems.append(f"{sheet.sheet_name}: {ws.max_row - 1} data rows, expected {n}")
        for i, scope in enumerate(model.scope):
            row = [c.value for c in ws[i + 2]]
            if row[1] != scope.tag_number:
                problems.append(f"{sheet.sheet_name} row {i + 2}: tag {row[1]!r}, expected {scope.tag_number!r}")
                break
            if scope.s_no is not None and (row[0] != scope.s_no or not isinstance(row[0], str)):
                problems.append(f"{sheet.sheet_name} row {i + 2}: S.NO {row[0]!r} is not the text {scope.s_no!r}")
                break
            source_values = row[n_scope:n_scope + n_src]
            # Guard 4 — nothing in a source column that the canonical layer does not hold.
            for src, got in zip(model.sources, source_values):
                want = sheet.cells[i][src.source_type].value
                if (got or None) != (want or None):
                    problems.append(f"{sheet.sheet_name} row {i + 2}: {src.source_type} holds {got!r}, canonical {want!r}")
            remark, final, status = row[n_scope + n_src], row[n_scope + n_src + 1], row[n_scope + n_src + 2]
            # Guard 1 — no verdict (and no FINAL) for a row with no source evidence at all. A duplicate note is
            # evidence: the source has rows, they just could not be reduced to one value.
            no_evidence = (all(v in (None, "") for v in source_values)
                           and not any(sheet.cells[i][s.source_type].note for s in model.sources))
            if no_evidence and (final not in (None, "") or status not in (None, "")):
                problems.append(f"{sheet.sheet_name} row {i + 2}: FINAL/STATUS set but no source holds a value")
            if status not in (None, "") and status not in ALLOWED_STATUS:
                problems.append(f"{sheet.sheet_name} row {i + 2}: STATUS {status!r} is not a comparison state")
            # Guard 5 — REMARKS / FINAL / STATUS are exactly the comparison result.
            for label, got, want in (("REMARKS", remark, model.column(model.remarks, sheet)[i]),
                                     ("FINAL", final, model.column(model.finals, sheet)[i]),
                                     ("STATUS", status, model.column(model.statuses, sheet)[i])):
                if (got or None) != (want or None):
                    problems.append(f"{sheet.sheet_name} row {i + 2}: {label} {got!r}, comparison result {want!r}")

    ws = wb[SUMMARY_SHEET]
    if [c.value for c in ws[1]] != model.summary_headers():
        problems.append(f"SUMMERY: headers {[c.value for c in ws[1]]} differ from the expected layout")
    else:
        for i, sheet in enumerate(model.sheets, start=2):
            if ws.cell(row=i, column=2).value != sheet.summary_description:
                problems.append(f"SUMMERY row {i}: description {ws.cell(row=i, column=2).value!r}")
            for j, src in enumerate(model.sources):
                got = ws.cell(row=i, column=3 + j).value
                if got != model.qty(sheet, src.source_type):
                    problems.append(f"SUMMERY row {i}: {src.summary_code} QTY {got!r} ≠ {model.qty(sheet, src.source_type)}")
            base = 3 + len(model.sources)
            for k, (h, want) in enumerate(zip(SUMMARY_RESULT_HEADERS, model.result_counts(sheet)), start=1):
                got = ws.cell(row=i, column=base + k).value
                if got != want:
                    problems.append(f"SUMMERY row {i}: {h} {got!r} ≠ {want!r}")
    return problems
