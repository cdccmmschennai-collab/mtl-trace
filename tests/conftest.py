import io
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.config import REPO_ROOT, Settings
from backend.main import create_app

REAL_SCOPE_FILE = REPO_ROOT / "input" / "references" / "5.MTL-DATA MISMATCH COMPARISION.xlsx"
REAL_DATASHEET_FILE = REPO_ROOT / "input" / "sources" / "datasheet" / "0904_4391-DATA SHEET PRIORITY-2_COMPLETED.xlsx"
REAL_SPIR_FILE = REPO_ROOT / "input" / "sources" / "spir" / "4391-SPIR DATA PRIORITY-2_COMPLETED.xlsx"
ORIGINAL_BACKUP_FILE = REPO_ROOT / "5.MTL-DATA MISMATCH COMPARISION - ORIGINAL BACKUP.xlsx"
# Header row of the real Data Sheet, in its real order (used to build realistic synthetic sources).
DS_HEADERS = ["S.NO", "TAG NUMBER", "EQUIPMENT DESCRIPTION", "DATASHEET REFERENCE-ORIGINAL", " TAG NUMBER AS PER DOC",
              "PAGE NO", "SIZE & RATING", "MAKE", "MODEL", "SERIAL NUMBER", "PART NUMBER", "REMARKS"]


def ds_row(tag, as_per_doc=None, make=None, model=None, serial=None, part=None, ref="DOC-1", page=1) -> list:
    return [None, tag, "EQUIP", ref, as_per_doc, page, None, make, model, serial, part, None]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = ["S.NO", "TAG NUMBER", "TAG DISCRIPTION", "EQUIPMENT DESCRIPTION", "SIZE & RATING"]


def make_workbook(sheets: dict[str, list[list]]) -> bytes:
    """Build an in-memory workbook: {sheet name: rows (first row is usually the header)}."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def scope_rows(n: int) -> list[list]:
    return [HEADERS] + [[f"{i:04d}", f"TAG-{i:03d}", f"DESC {i}", "EQUIP", None] for i in range(1, n + 1)]


@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    return tmp_path / "MTL_DATA"


@pytest.fixture
def app(data_dir: Path):
    app = create_app(Settings(data_dir=data_dir))
    yield app
    app.state.db.dispose()


@pytest.fixture
def client(app) -> TestClient:
    with TestClient(app) as c:
        yield c


@pytest.fixture
def real_scope_bytes() -> bytes:
    if not REAL_SCOPE_FILE.exists():
        pytest.skip("real scope workbook not present")
    return REAL_SCOPE_FILE.read_bytes()


@pytest.fixture
def real_datasheet_bytes() -> bytes:
    if not REAL_DATASHEET_FILE.exists():
        pytest.skip("real Data Sheet workbook not present")
    return REAL_DATASHEET_FILE.read_bytes()


@pytest.fixture
def real_spir_bytes() -> bytes:
    if not REAL_SPIR_FILE.exists():
        pytest.skip("real SPIR workbook not present")
    return REAL_SPIR_FILE.read_bytes()


def without_column(data: bytes, sheet: str, header: str) -> bytes:
    """A copy of a real workbook with the column whose header reads `header` removed (found by header text)."""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb[sheet]
    [col] = [c.column for c in ws[1] if c.value is not None and str(c.value).strip() == header]
    ws.delete_cols(col)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def with_extra_row(data: bytes, sheet: str, values: dict[str, object]) -> bytes:
    """A copy of a real workbook with one row appended, values placed by header text."""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb[sheet]
    cols = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    r = ws.max_row + 1
    for header, v in values.items():
        ws.cell(row=r, column=cols[header], value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client: TestClient, path: str, data: bytes, name: str = "scope.xlsx", **form):
    return client.post(path, files={"file": (name, data, XLSX)}, data=form)


def locked_milestone(client: TestClient, scope_bytes: bytes, code: str = "MTL-2026-001") -> dict:
    m = create(client, code)
    assert upload(client, f"/api/milestones/{m['id']}/scope", scope_bytes).status_code == 201
    r = client.post(f"/api/milestones/{m['id']}/scope/lock")
    assert r.status_code == 200, r.text
    return r.json()


def new_run(client: TestClient, milestone_id: int) -> dict:
    r = client.post(f"/api/milestones/{milestone_id}/runs", json={})
    assert r.status_code == 201, r.text
    return r.json()


def create(client: TestClient, code: str = "MTL-2026-001", **kw) -> dict:
    r = client.post("/api/milestones", json={"milestone_code": code, **kw})
    assert r.status_code == 201, r.text
    return r.json()
