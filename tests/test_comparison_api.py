"""Integration tests (Phase 1D): real 717-tag scope + real Data Sheet + real SPIR → consolidation → comparison
→ comparison workbook, over HTTP."""

import hashlib
import io
import re
import sqlite3

import openpyxl
import pytest

from backend.domain.sources import ATTRIBUTE_INFO, ATTRIBUTES
from tests.conftest import (
    REAL_DATASHEET_FILE,
    REAL_SPIR_FILE,
    locked_milestone,
    make_workbook,
    new_run,
    scope_rows,
    upload,
    without_column,
)
from tests.test_consolidation_api import consolidate, ds_file, ds_row

STATES = ("MATCH", "MISMATCH", "REVIEW REQUIRED")


def compare(client, run_id: int) -> dict:
    r = client.post(f"/api/runs/{run_id}/compare")
    assert r.status_code == 201, r.text
    return r.json()


def workbook(client, run_id: int, cmp: dict):
    art = cmp["output"]
    r = client.get(f"/api/runs/{run_id}/outputs/{art['id']}/download")
    assert r.status_code == 200 and hashlib.sha256(r.content).hexdigest() == art["sha256"]
    assert art["kind"] == "COMPARISON_WORKBOOK" and art["file_name"].endswith("RUN-001.xlsx")
    return openpyxl.load_workbook(io.BytesIO(r.content))


@pytest.fixture
def real_compared(client, real_scope_bytes, real_datasheet_bytes, real_spir_bytes):
    m = locked_milestone(client, real_scope_bytes)
    run = new_run(client, m["id"])
    upload(client, f"/api/runs/{run['id']}/sources", real_datasheet_bytes, name=REAL_DATASHEET_FILE.name)
    upload(client, f"/api/runs/{run['id']}/sources", real_spir_bytes, name=REAL_SPIR_FILE.name)
    consolidate(client, run["id"])
    cmp = compare(client, run["id"])
    return m, run, cmp, workbook(client, run["id"], cmp)


def same_maker_possibly(a: str, b: str) -> bool:
    """Manufacturer-name variant, restated: same words ignoring punctuation/spacing, one name's words all inside
    the other's, or equal but for a truncated last word."""
    x, y = re.findall(r"[0-9a-z]+", a.lower()), re.findall(r"[0-9a-z]+", b.lower())
    return ("".join(x) == "".join(y) or set(x).issubset(y) or set(y).issubset(x)
            or (len(x) == len(y) > 1 and x[:-1] == y[:-1] and (x[-1].startswith(y[-1]) or y[-1].startswith(x[-1]))))


def expected(attribute: str, milestone_tag: str, values: list) -> tuple:
    """An independent restatement of the rules, applied to the workbook's own source columns (priority order)."""
    vals = [str(x).strip() for x in values if x is not None and str(x).strip()]
    if not vals:
        return None, None
    distinct = {x.casefold() for x in vals}
    if attribute == "TAG_NUMBER":
        if len(distinct) > 1:
            return milestone_tag, "MISMATCH"
        return milestone_tag, "MATCH" if distinct == {milestone_tag.strip().casefold()} else "REVIEW REQUIRED"
    if attribute == "MAKE" and len(distinct) > 1 and all(
            same_maker_possibly(a, b) for i, a in enumerate(vals) for b in vals[i + 1:] if a.casefold() != b.casefold()):
        return None, "REVIEW REQUIRED"
    return vals[0], "MISMATCH" if len(distinct) > 1 else "MATCH"


def test_real_comparison_counts(real_compared):
    m, _, cmp, _ = real_compared
    n = m["tag_count"]
    assert cmp["scope_tag_count"] == n == 717
    assert [t["code"] for t in cmp["source_types"]] == ["MDS", "SPIR"]
    assert [t["rank"] for t in cmp["source_types"]] == [2, 3]
    assert sum(cmp["totals"].values()) + cmp["no_source_value"] == n * len(ATTRIBUTES)
    for a in cmp["attributes"]:
        assert sum(a["counts"].values()) + a["no_source_value"] == n
        assert a["provided_by"] == ["MDS", "SPIR"]
    by = {a["code"]: a for a in cmp["attributes"]}
    # Regression figures for the real Data Sheet + SPIR (rules v1).
    assert cmp["totals"] == {"MATCH": 1852, "MISMATCH": 63, "REVIEW REQUIRED": 77}
    assert by["TAG_NUMBER"]["counts"] == {"MATCH": 410, "MISMATCH": 11, "REVIEW REQUIRED": 23}
    assert by["MAKE"]["counts"] == {"MATCH": 394, "MISMATCH": 13, "REVIEW REQUIRED": 54}   # 54 naming variants
    assert by["MODEL"]["counts"] == {"MATCH": 347, "MISMATCH": 39, "REVIEW REQUIRED": 0}
    assert by["SERIAL_NUMBER"]["counts"] == {"MATCH": 331, "MISMATCH": 0, "REVIEW REQUIRED": 0}
    assert by["PART_NUMBER"]["counts"] == {"MATCH": 370, "MISMATCH": 0, "REVIEW REQUIRED": 0}


def test_real_workbook_structure(real_compared):
    _, _, cmp, wb = real_compared
    assert wb.sheetnames == ["SUMMERY", *(ATTRIBUTE_INFO[a].sheet_name for a in ATTRIBUTES)]
    for a in ATTRIBUTES:
        info = ATTRIBUTE_INFO[a]
        assert [c.value for c in wb[info.sheet_name][1]] == [
            "S.NO", "TAG NUMBER", "TAG DISCRIPTION", "EQUIPMENT DESCRIPTION", "SIZE & RATING",
            f"{info.column_label} IN DATA SHEET (MDS)", f"{info.column_label} IN SPIR (MIR)",
            "REMARKS", info.final_label, "STATUS"]
    assert [s["headers"] for s in cmp["workbook"]["sheets"]] == [[c.value for c in wb[n][1]] for n in wb.sheetnames[1:]]
    s = wb["SUMMERY"]
    assert [c.value for c in s[1]] == ["SL NO", "DESCRIPTION", "MDS QTY", "MIR QTY", "REMARKS", *STATES]
    for r, a in enumerate(cmp["attributes"], start=2):
        assert [s.cell(row=r, column=c).value for c in (6, 7, 8)] == [a["counts"][h] for h in STATES]
    assert s["C8"].value == 717


def test_every_real_row_follows_the_rules(real_compared):
    """Oracle: recompute FINAL and STATUS for all 3,585 rows from the workbook's own source columns."""
    _, _, _, wb = real_compared
    for a in ATTRIBUTES:
        ws = wb[ATTRIBUTE_INFO[a].sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            tag, mds, spir, remarks, final, status = row[1], row[5], row[6], row[7], row[8], row[9]
            assert (final, status) == expected(a.value, tag, [mds, spir]), (a.value, tag, mds, spir, final, status)
            assert status in (None, *STATES)
            assert (remarks is None) == (status in (None, "MATCH")), (a.value, tag)   # every non-MATCH is explained


def test_business_examples_on_real_rows(real_compared):
    _, _, _, wb = real_compared
    rows = {r[1]: r for r in wb["TAG NUMBER COMPARISION"].iter_rows(min_row=2, values_only=True)}
    # Sources disagree with each other → MISMATCH, FINAL TAG = milestone tag, source values kept.
    assert rows["68-FIT-1987"][5:10] == ("68-FIT-1987", "68-FT-1987", rows["68-FIT-1987"][7], "68-FIT-1987", "MISMATCH")
    # One source value differs from the milestone tag, no source-to-source disagreement → REVIEW REQUIRED.
    r = rows["68-SI-01-10-1805"]
    assert (r[5], r[6], r[8], r[9]) == (None, "68-SI-10-1805", "68-SI-01-10-1805", "REVIEW REQUIRED")
    make = {r[1]: r for r in wb["MAKE COMPARISION"].iter_rows(min_row=2, values_only=True)}
    r = make["69-FD-10-141"]                               # manufacturer-name variant — REVIEW REQUIRED, FINAL open
    assert (r[5], r[6], r[7], r[8], r[9]) == ("TELEDYNE", "SIMTRONICS TELEDYNE",
                                              "MAKE values may represent the same manufacturer; human review required.",
                                              None, "REVIEW REQUIRED")
    serial = {r[1]: r for r in wb["SERIAL NUMBER COMPARISION"].iter_rows(min_row=2, values_only=True)}
    r = serial["6834-P-64A"]                               # MDS empty does not override SPIR
    assert (r[5], r[6], r[8], r[9]) == (None, "665436", "665436", "MATCH")


def test_real_rows_without_any_source_value_have_no_verdict(real_compared):
    _, _, cmp, wb = real_compared
    blank = 0
    for a in ATTRIBUTES:
        for row in wb[ATTRIBUTE_INFO[a].sheet_name].iter_rows(min_row=2, values_only=True):
            if row[5] is None and row[6] is None:
                blank += 1
                assert row[7:10] == (None, None, None)
    assert blank == cmp["no_source_value"]
    everything = {c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row}
    assert not everything & {"PENDING", "NOT AVAILABLE", "UNKNOWN", "RESOLVED", "PENDING/NOT AVAILABLE"}


def test_comparison_is_persisted_once_and_immutable(client, real_compared):
    _, run, cmp, _ = real_compared
    r = client.post(f"/api/runs/{run['id']}/compare")
    assert r.status_code == 409 and "new run" in r.json()["detail"]
    assert client.get(f"/api/runs/{run['id']}/comparison").json() == cmp
    con = sqlite3.connect(client.app.state.db.path, isolation_level=None)
    assert con.execute("select count(*) from comparison_result where run_id = ?", (run["id"],)).fetchone()[0] == 717 * 5
    assert {r[0] for r in con.execute("select distinct status from comparison_result")} == {None, *STATES}
    for sql in ("UPDATE comparison_result SET status = 'MATCH'", "DELETE FROM comparison_result",
                "UPDATE run_comparison SET summary = '{}'"):
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            con.execute(sql)


# ---------------------------------------------------------------- workflow, dynamic columns, missing fields

def test_compare_requires_consolidation(client):
    m = locked_milestone(client, make_workbook({"S": scope_rows(3)}))
    run = new_run(client, m["id"])
    upload(client, f"/api/runs/{run['id']}/sources", ds_file([ds_row("TAG-001", make="A")]))
    r = client.post(f"/api/runs/{run['id']}/compare")
    assert r.status_code == 409 and "Consolidate" in r.json()["detail"]
    assert client.get(f"/api/runs/{run['id']}/comparison").status_code == 404
    assert client.post("/api/runs/999/compare").status_code == 404


def test_source_columns_are_the_participating_types_only(client, real_scope_bytes, real_datasheet_bytes):
    m = locked_milestone(client, real_scope_bytes)
    run = new_run(client, m["id"])
    upload(client, f"/api/runs/{run['id']}/sources", real_datasheet_bytes, name=REAL_DATASHEET_FILE.name)
    consolidate(client, run["id"])
    cmp = compare(client, run["id"])
    wb = workbook(client, run["id"], cmp)
    assert [t["code"] for t in cmp["source_types"]] == ["MDS"]
    assert [c.value for c in wb["MAKE COMPARISION"][1]][5:] == ["MAKE IN DATA SHEET (MDS)", "REMARKS", "FINAL MAKE", "STATUS"]
    # Data Sheet alone: one source never disagrees with itself; TAG rows can still need review.
    assert cmp["totals"]["MISMATCH"] == 0
    by = {a["code"]: a for a in cmp["attributes"]}
    assert by["TAG_NUMBER"]["counts"]["REVIEW REQUIRED"] == 14          # the 14 Data Sheet tags that differ
    assert by["MAKE"]["counts"]["MATCH"] == 124 and by["PART_NUMBER"]["no_source_value"] == 717


def test_attribute_a_source_does_not_provide_creates_no_mismatch(client, real_scope_bytes, real_datasheet_bytes,
                                                                    real_spir_bytes):
    m = locked_milestone(client, real_scope_bytes)
    run = new_run(client, m["id"])
    upload(client, f"/api/runs/{run['id']}/sources", real_datasheet_bytes, name=REAL_DATASHEET_FILE.name)
    upload(client, f"/api/runs/{run['id']}/sources", without_column(real_spir_bytes, "SPIR DATA", "SERIAL NUMBER"),
           name="spir-no-serial.xlsx")
    consolidate(client, run["id"])
    cmp = compare(client, run["id"])
    serial = next(a for a in cmp["attributes"] if a["code"] == "SERIAL_NUMBER")
    assert serial["provided_by"] == ["MDS"]
    assert serial["counts"] == {"MATCH": 0, "MISMATCH": 0, "REVIEW REQUIRED": 0} and serial["no_source_value"] == 717
    ws = workbook(client, run["id"], cmp)["SERIAL NUMBER COMPARISION"]
    assert [c.value for c in ws[1]][5:7] == ["SL/NO IN DATA SHEET (MDS)", "SL/NO IN SPIR (MIR)"]   # column kept, empty
    assert all(r[6] is None and r[8] is None and r[9] is None for r in ws.iter_rows(min_row=2, values_only=True))


def test_a_third_source_type_participates_automatically(client):
    """Data Sheet + SPIR + an MXS document (identified by keyword): MXS gets a column in priority order and
    takes part in the comparison with no code change."""
    m = locked_milestone(client, make_workbook({"S": scope_rows(3)}))
    run = new_run(client, m["id"])
    upload(client, f"/api/runs/{run['id']}/sources",
           ds_file([ds_row("TAG-001", as_per_doc="TAG-001", make="ACME"), ds_row("TAG-002", make=None)]))
    upload(client, f"/api/runs/{run['id']}/sources", make_workbook({"MXS DATA": [
        ["TAG NUMBER", "MXS TAG NUMBER", "MAKE", "MODEL"],
        ["TAG-001", "TAG-001X", "ACME", "M-1"], ["TAG-002", None, "MXS MAKE", None]]}), name="mxs.xlsx")
    upload(client, f"/api/runs/{run['id']}/sources", make_workbook({"SPIR DATA": [
        ["TAG NUMBER", "SPIR TAG NUMBER", "MAKE"], ["TAG-001", "TAG-001", "acme "]]}), name="spir.xlsx")
    consolidate(client, run["id"])
    cmp = compare(client, run["id"])
    assert [t["code"] for t in cmp["source_types"]] == ["MDS", "SPIR", "MXS"]
    by = {a["code"]: a for a in cmp["attributes"]}
    assert by["MODEL"]["provided_by"] == ["MDS", "MXS"]                  # SPIR has no MODEL column
    wb = workbook(client, run["id"], cmp)
    ws = wb["MAKE COMPARISION"]
    assert [c.value for c in ws[1]][5:] == ["MAKE IN DATA SHEET (MDS)", "MAKE IN SPIR (MIR)", "MAKE IN CROSS SECTION (MXS)",
                                           "REMARKS", "FINAL MAKE", "STATUS"]
    rows = {r[1]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert (rows["TAG-001"][9], rows["TAG-001"][10]) == ("ACME", "MATCH")          # ACME = acme (case, trim)
    assert (rows["TAG-002"][9], rows["TAG-002"][10]) == ("MXS MAKE", "MATCH")      # lower-priority fills the gap
    assert rows["TAG-003"][8:] == (None, None, None)
    tag = {r[1]: r for r in wb["TAG NUMBER COMPARISION"].iter_rows(min_row=2, values_only=True)}
    assert (tag["TAG-001"][9], tag["TAG-001"][10]) == ("TAG-001", "MISMATCH")      # MDS/SPIR ≠ MXS
    assert [c.value for c in wb["SUMMERY"][1]][2:5] == ["MDS QTY", "MIR QTY", "MXS QTY"]
