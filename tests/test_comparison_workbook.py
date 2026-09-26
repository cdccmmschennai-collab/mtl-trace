"""Unit tests: dynamic generation of the comparison workbook structure (backend/output)."""

import io
from pathlib import Path

import openpyxl
import pytest

from backend.config import REPO_ROOT
from backend.domain.sources import ATTRIBUTE_INFO, ATTRIBUTES, SOURCE_TYPES, ordered_source_types
from backend.output.comparison_workbook import (
    SCOPE_HEADERS,
    AttributeSheet,
    ScopeRow,
    SourceCell,
    SourceColumn,
    WorkbookModel,
    WorkbookValidationError,
    build_workbook,
    validate_workbook,
)
from tests.conftest import REAL_SCOPE_FILE

AUTOMATED_FILE = REPO_ROOT / "5.MTL-DATA MISMATCH COMPARISION - AUTOMATED.xlsx"
TEMPLATE_SOURCES = ["ASSET", "MDS", "SPIR", "MXS", "MOM", "MTC"]      # the six columns of the existing workbook


def model(source_types: list[str], n: int = 3, values: dict | None = None, columns=None) -> WorkbookModel:
    """values: {(row index, attribute, source type): value}."""
    values = values or {}
    cols = columns or [SourceColumn(st, SOURCE_TYPES[st].label, SOURCE_TYPES[st].summary_code)
                       for st in ordered_source_types(source_types)]
    sheets = []
    for a in ATTRIBUTES:
        info = ATTRIBUTE_INFO[a]
        sheets.append(AttributeSheet(a.value, info.sheet_name, info.column_label, info.final_label,
                                     info.summary_description,
                                     [{c.source_type: SourceCell(values.get((i, a.value, c.source_type)))
                                       for c in cols} for i in range(n)]))
    scope = [ScopeRow(f"{i + 1:04d}", f"TAG-{i + 1:03d}", "DESC", "EQUIP", None) for i in range(n)]
    return WorkbookModel(scope=scope, sources=cols, sheets=sheets)


def load(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


def headers(ws) -> list:
    return [c.value for c in ws[1]]


# ---------------------------------------------------------------- template fidelity

def test_sheet_names_match_the_documented_structure():
    wb = load(build_workbook(model(["MDS"])))
    assert wb.sheetnames == ["SUMMERY", "TAG NUMBER COMPARISION", "MAKE COMPARISION", "MODEL COMPARISION",
                             "SERIAL NUMBER COMPARISION", "PART NUMBER COMPARISION"]
    if AUTOMATED_FILE.exists():                                  # layout reference (structure only)
        assert wb.sheetnames == openpyxl.load_workbook(AUTOMATED_FILE, read_only=True).sheetnames


def test_template_scope_headers_and_summary_headers_are_reproduced():
    if not REAL_SCOPE_FILE.exists():
        pytest.skip("working file not present")
    template = openpyxl.load_workbook(REAL_SCOPE_FILE, read_only=True)
    ours = load(build_workbook(model(TEMPLATE_SOURCES)))
    theirs = [c.value for c in next(template["SUMMERY"].iter_rows(max_row=1))]
    mine = headers(ours["SUMMERY"])
    # With the template's six sources the dynamic SUMMERY header row carries the template's headers verbatim,
    # in the template's positions — except that source columns follow the documented priority sequence
    # (PRD §13.2: … MXS · MTC · MOM), whereas the template happens to place MOM before MTC, and that the
    # template's legacy "PENDING/NOT AVAILABLE" result header is not reproduced (not a comparison state).
    assert mine == [*theirs[:6], "MTC QTY", "MOM QTY", *[h for h in theirs[8:] if h != "PENDING/NOT AVAILABLE"]]
    assert theirs[6:8] == ["MOM QTY", "MTC QTY"] and "PENDING/NOT AVAILABLE" in theirs
    assert mine[-3:] == ["MATCH", "MISMATCH", "REVIEW REQUIRED"]
    assert [ours["SUMMERY"].cell(row=r, column=2).value for r in range(2, 7)] == \
        [template["SUMMERY"].cell(row=r, column=2).value for r in range(2, 7)]
    assert headers(ours["TAG NUMBER COMPARISION"])[:5] == \
        [c.value for c in next(template["TAG NUMBER COMPARISION"].iter_rows(max_row=1))][:5]


def test_result_columns_follow_the_mental_model_reference():
    ours = load(build_workbook(model(["MDS"])))
    ref = openpyxl.load_workbook(AUTOMATED_FILE, read_only=True) if AUTOMATED_FILE.exists() else None
    for a in ATTRIBUTES:
        name = ATTRIBUTE_INFO[a].sheet_name
        h = headers(ours[name])
        assert h[-3:] == ["REMARKS", ATTRIBUTE_INFO[a].final_label, "STATUS"]
        if ref is not None:
            assert [c.value for c in next(ref[name].iter_rows(max_row=1))][-3:] == h[-3:]


def test_formatting_matches_the_reference():
    ws = load(build_workbook(model(["MDS", "SPIR"])))["MAKE COMPARISION"]
    for col in range(1, 6):
        c = ws.cell(row=1, column=col)
        assert c.fill.fgColor.rgb == "FF7030A0" and c.font.bold
    for col in range(6, ws.max_column + 1):
        assert ws.cell(row=1, column=col).fill.fgColor.rgb == "FFFFFF00"
    assert ws.freeze_panes == "A2"
    assert ws["A2"].value == "0001" and ws["A2"].data_type == "s"          # S.NO written as text


# ---------------------------------------------------------------- dynamic source columns

@pytest.mark.parametrize("types, expected", [
    (["MDS"], ["DATA SHEET (MDS)"]),
    (["SPIR", "MDS"], ["DATA SHEET (MDS)", "SPIR (MIR)"]),
    (["MDS", "GA", "SPIR"], ["DATA SHEET (MDS)", "SPIR (MIR)", "GA DOCUMENT"]),
    (["MOM", "MTC", "MXS", "MDS"], ["DATA SHEET (MDS)", "CROSS SECTION (MXS)", "TEST CERTIFICATE (MTC)", "O&M MANUAL (MOM)"]),
])
def test_source_columns_are_exactly_the_participating_types_in_display_order(types, expected):
    wb = load(build_workbook(model(types)))
    for a in ATTRIBUTES:
        info = ATTRIBUTE_INFO[a]
        h = headers(wb[info.sheet_name])
        assert h[5:-3] == [f"{info.column_label} IN {label}" for label in expected]
    codes = [SOURCE_TYPES[t].summary_code for t in ordered_source_types(types)]
    assert headers(wb["SUMMERY"])[2:2 + len(types)] == [f"{c} QTY" for c in codes]


def test_ga_appears_only_when_ga_participates():
    without = load(build_workbook(model(["MDS", "SPIR"])))
    with_ga = load(build_workbook(model(["MDS", "SPIR", "GA"])))
    for ws in without.worksheets:
        assert not any("GA" in str(v).split() for v in headers(ws)), ws.title
    assert headers(with_ga["MAKE COMPARISION"])[7] == "MAKE IN GA DOCUMENT"
    assert "GA QTY" in headers(with_ga["SUMMERY"])


def test_a_new_source_type_needs_no_code_change():
    col = SourceColumn("NEWDOC", "NEW DOCUMENT TYPE (NDT)", "NDT")
    m = model([], columns=[SourceColumn("MDS", "DATA SHEET (MDS)", "MDS"), col],
              values={(0, "MAKE", "NEWDOC"): "ACME"})
    wb = load(build_workbook(m))
    assert headers(wb["MAKE COMPARISION"])[6] == "MAKE IN NEW DOCUMENT TYPE (NDT)"
    assert wb["MAKE COMPARISION"]["G2"].value == "ACME"
    assert "NDT QTY" in headers(wb["SUMMERY"])


def test_writer_hard_codes_no_source_column():
    text = (REPO_ROOT / "backend" / "output" / "comparison_workbook.py").read_text(encoding="utf-8")
    for code, info in SOURCE_TYPES.items():
        assert info.label not in text and f'"{code}"' not in text, code


# ---------------------------------------------------------------- values and summary

def test_values_written_verbatim_and_summary_quantities_computed():
    values = {(0, "MAKE", "MDS"): "TOPSAFE CO., LTD", (1, "MAKE", "MDS"): " Honeywell",
              (0, "MAKE", "SPIR"): "=NOT A FORMULA", (2, "MODEL", "SPIR"): "M-1"}
    wb = load(build_workbook(model(["MDS", "SPIR"], values=values)))
    ws = wb["MAKE COMPARISION"]
    assert (ws["F2"].value, ws["F3"].value, ws["G2"].value) == ("TOPSAFE CO., LTD", " Honeywell", "=NOT A FORMULA")
    assert ws["G2"].data_type == "s"
    s = wb["SUMMERY"]
    assert [s.cell(row=3, column=c).value for c in (3, 4)] == [2, 1]          # MAKE: MDS 2, MIR 1
    assert [s.cell(row=4, column=c).value for c in (3, 4)] == [0, 1]          # MODEL
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for row in s.iter_rows() for c in row)


def test_phase_1c_writes_no_comparison_result():
    wb = load(build_workbook(model(["MDS", "SPIR"], values={(0, "MAKE", "MDS"): "A", (0, "MAKE", "SPIR"): "B"})))
    for a in ATTRIBUTES:
        ws = wb[ATTRIBUTE_INFO[a].sheet_name]
        for col in range(ws.max_column - 2, ws.max_column + 1):              # REMARKS, FINAL, STATUS
            assert all(ws.cell(row=r, column=col).value is None for r in range(2, ws.max_row + 1))
    s = wb["SUMMERY"]
    for r in range(2, 7):
        assert all(s.cell(row=r, column=c).value is None for c in range(5, s.max_column + 1))
    everything = {c.value for ws in wb.worksheets for row in ws.iter_rows(min_row=2) for c in row}
    assert not everything & {"MATCH", "MISMATCH", "REVIEW REQUIRED", "PENDING", "NOT AVAILABLE", "RESOLVED"}


def test_comparison_result_is_written_and_counted():
    m = model(["MDS", "SPIR"], values={(0, "MAKE", "MDS"): "A", (0, "MAKE", "SPIR"): "B", (1, "MAKE", "SPIR"): "C"})
    m.statuses = {"MAKE": ["MISMATCH", "MATCH", None]}
    m.finals = {"MAKE": ["A", "C", None]}
    m.remarks = {"MAKE": ["Sources disagree", None, None]}
    data = build_workbook(m)
    wb = load(data)
    ws = wb["MAKE COMPARISION"]
    assert [ws.cell(row=r, column=c).value for r in (2, 3, 4) for c in (8, 9, 10)] == \
        ["Sources disagree", "A", "MISMATCH", None, "C", "MATCH", None, None, None]
    assert [wb["SUMMERY"].cell(row=3, column=c).value for c in (6, 7, 8)] == [1, 1, 0]
    assert [wb["SUMMERY"].cell(row=2, column=c).value for c in (6, 7, 8)] == [None, None, None]   # TAG: no result
    assert any("STATUS" in p for p in validate_workbook(_tamper(data, "MAKE COMPARISION", "J3", "MISMATCH"), m))


def test_model_with_a_non_state_status_is_refused():
    m = model(["MDS"], values={(0, "MAKE", "MDS"): "A"})
    m.statuses = {"MAKE": ["PENDING", None, None]}
    with pytest.raises(ValueError, match="not comparison states"):
        build_workbook(m)


def test_total_tag_count_row():
    s = load(build_workbook(model(["MDS"], n=4)))["SUMMERY"]
    assert (s["B8"].value, s["C8"].value) == ("TOTAL EQUIPMENT / TAG COUNT", 4)


# ---------------------------------------------------------------- output guards

def _tamper(data: bytes, sheet: str, coord: str, value) -> bytes:
    wb = load(data)
    wb[sheet][coord] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_guard_value_not_held_by_canonical_layer():
    m = model(["MDS"])
    bad = _tamper(build_workbook(m), "MAKE COMPARISION", "F2", "INVENTED")
    assert any("canonical" in p for p in validate_workbook(bad, m))


def test_guard_verdict_on_row_without_source_values():
    m = model(["MDS"])
    bad = _tamper(build_workbook(m), "MAKE COMPARISION", "I2", "MATCH")     # STATUS on an all-empty row
    assert any("no source holds a value" in p for p in validate_workbook(bad, m))


def test_guard_non_state_status():
    m = model(["MDS"], values={(0, "MAKE", "MDS"): "A"})
    bad = _tamper(build_workbook(m), "MAKE COMPARISION", "I2", "PENDING")
    assert any("not a comparison state" in p for p in validate_workbook(bad, m))


def test_guard_row_order_and_extra_columns():
    m = model(["MDS"])
    assert any("tag" in p for p in validate_workbook(_tamper(build_workbook(m), "MAKE COMPARISION", "B2", "X"), m))
    assert any("headers" in p for p in validate_workbook(
        _tamper(build_workbook(m), "MAKE COMPARISION", "J1", "MAKE IN GA DOCUMENT"), m))


def test_model_with_wrong_row_count_is_refused():
    m = model(["MDS"])
    m.sheets[0].cells.pop()
    with pytest.raises(ValueError):
        build_workbook(m)


def test_validation_error_carries_problems():
    err = WorkbookValidationError(["a", "b"])
    assert err.problems == ["a", "b"] and "a" in str(err)
