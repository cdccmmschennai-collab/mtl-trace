"""Integration tests (Phase 1C): consolidation over HTTP, the canonical dataset, run isolation, generated outputs.

Test A — Data Sheet only (real 717-tag milestone).
Test B — Data Sheet + CONTROLLED SPIR FIXTURE (tests/fixtures/second_source.py — replaces the real SPIR
         adapter for these tests only, to exercise duplicates / out-of-scope in a second source).
Test C — an earlier run is unchanged when a later run adds a source.
Test D — real Data Sheet + real SPIR (Phase 1D-1), incl. a missing column and an out-of-scope tag.
"""

import hashlib
import io
import json
import sqlite3

import openpyxl
import pytest

from backend.domain.sources import ATTRIBUTE_INFO, ATTRIBUTES
from backend.extraction import registry
from tests.conftest import (
    DS_HEADERS,
    ORIGINAL_BACKUP_FILE,
    REAL_DATASHEET_FILE,
    REAL_SCOPE_FILE,
    REAL_SPIR_FILE,
    ds_row,
    locked_milestone,
    make_workbook,
    new_run,
    scope_rows,
    upload,
    with_extra_row,
    without_column,
)
from tests.fixtures.second_source import FIXTURE_SPIR_ADAPTER, fixture_row, fixture_workbook
from tests.test_datasheet_adapter import KNOWN_BACKUP_DEVIATIONS
from tests.test_spir_adapter import KNOWN_SPIR_BACKUP_DEVIATIONS, independent_spir


@pytest.fixture
def spir_fixture(monkeypatch):
    """Register the controlled second-source fixture for this test only."""
    monkeypatch.setitem(registry.ADAPTERS, "SPIR", FIXTURE_SPIR_ADAPTER)


def db(client) -> sqlite3.Connection:
    return sqlite3.connect(client.app.state.db.path, isolation_level=None)


def ds_file(rows: list[list]) -> bytes:
    return make_workbook({"DATASHEET": [DS_HEADERS] + rows})


def consolidate(client, run_id: int) -> dict:
    r = client.post(f"/api/runs/{run_id}/consolidate")
    assert r.status_code == 201, r.text
    return r.json()


def canonical(client, run_id: int, **params) -> dict:
    q = "&".join(f"{k}={v}" for k, v in {"limit": 1000, **params}.items())
    r = client.get(f"/api/runs/{run_id}/canonical?{q}")
    assert r.status_code == 200, r.text
    return r.json()


def workbook_of(client, data_dir, cons: dict):
    [art] = [o for o in cons["outputs"] if o["kind"] == "CONSOLIDATED_WORKBOOK"]
    data = (data_dir / art["stored_path"]).read_bytes()
    assert hashlib.sha256(data).hexdigest() == art["sha256"]
    return openpyxl.load_workbook(io.BytesIO(data)), art


def scope_tags_of(scope_bytes: bytes) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(scope_bytes), read_only=True)
    rows = list(wb["TAG NUMBER COMPARISION"].iter_rows(min_row=2, values_only=True))
    return [r[1] for r in rows if r[1]]


# ================================================================ Test A — real Data Sheet only

@pytest.fixture
def real_run(client, real_scope_bytes, real_datasheet_bytes):
    m = locked_milestone(client, real_scope_bytes)
    run = new_run(client, m["id"])
    src = upload(client, f"/api/runs/{run['id']}/sources", real_datasheet_bytes, name=REAL_DATASHEET_FILE.name).json()
    return m, run, src, consolidate(client, run["id"])


def test_real_717_tag_consolidation(real_run, real_scope_bytes):
    m, run, src, cons = real_run
    scope_count = len(scope_tags_of(real_scope_bytes))                 # derived from the scope file
    assert cons["scope_tag_count"] == cons["canonical_row_count"] == m["tag_count"] == scope_count
    assert cons["canonical_cell_count"] == scope_count * len(ATTRIBUTES) * 1
    assert [t["code"] for t in cons["source_types"]] == ["MDS"]
    assert {t["code"] for t in cons["not_in_run"]} == {"ASSET", "SPIR", "GA", "MXS", "MTC", "MOM"}

    [cov] = cons["coverage"]
    # Derived relationships with the Phase 1B source processing result …
    assert cov["rows"] == src["row_count"] and cov["unique_tags"] == src["unique_tag_count"]
    assert cov["in_scope_tags"] == src["matched_tag_count"]
    assert cov["scope_tags_without_row"] == scope_count - cov["in_scope_tags"]
    for a in ATTRIBUTES:
        st = cov["attributes"][a.value]
        assert st["present"] == src["attribute_counts"][a.value]
        assert st["present"] + st["absent_value"] + st["absent_tag"] + st["duplicate"] == scope_count
        assert st["absent_tag"] == cov["scope_tags_without_row"]
    # … and the verified golden regression figures.
    assert (cov["rows"], cov["unique_tags"], cov["duplicate_tag_count"], cov["in_scope_tags"],
            cov["out_of_scope_tag_count"], cov["scope_tags_without_row"]) == (436, 436, 0, 436, 0, 281)
    assert {a: v["present"] for a, v in cov["attributes"].items()} == \
        {"TAG_NUMBER": 357, "MAKE": 124, "MODEL": 98, "SERIAL_NUMBER": 0, "PART_NUMBER": 0}
    assert cov["attributes"]["TAG_NUMBER"]["absent_value"] == 79
    assert (cov["tag_as_per_doc_equal"], cov["tag_as_per_doc_differs"]) == (343, 14)
    assert {"milestone_tag": "68-SP-051", "tag_as_per_doc": "SP-123"} in cov["tag_differences"]


def test_real_canonical_rows_are_the_scope_in_order(client, real_run, real_scope_bytes):
    _, run, _, _ = real_run
    page = canonical(client, run["id"])
    assert page["total"] == page["filtered_total"] == len(page["rows"]) == 717
    assert [r["tag_number"] for r in page["rows"]] == scope_tags_of(real_scope_bytes)
    assert page["rows"][0]["s_no"] == "0001"
    assert canonical(client, run["id"], filter="no_source_row")["filtered_total"] == 281
    assert canonical(client, run["id"], filter="tag_differs")["filtered_total"] == 14


def test_real_values_preserved_with_provenance(client, real_run):
    _, run, _, _ = real_run
    [row] = canonical(client, run["id"], search="68-SP-015")["rows"]
    make = row["cells"]["MAKE"]["MDS"]
    assert (make["state"], make["value"]) == ("PRESENT", "TOPSAFE CO., LTD")     # not "TOPSAFE CO LTD"
    [p] = make["provenance"]
    assert (p["document"], p["sheet"], p["cell"]) == (REAL_DATASHEET_FILE.name, "DATASHEET", "J428")
    assert row["source_rows"] == {"MDS": "ROW"}
    [row] = canonical(client, run["id"], search="68-SP-051")["rows"]
    tag = row["cells"]["TAG_NUMBER"]["MDS"]
    assert (row["tag_number"], tag["value"], tag["provenance"][0]["source_tag"]) == ("68-SP-051", "SP-123", "68-SP-051")


def test_real_generated_workbook_structure(client, data_dir, real_run):
    _, run, _, cons = real_run
    wb, art = workbook_of(client, data_dir, cons)
    assert art["stored_path"].startswith("milestones/MTL-2026-001/outputs/run-001/")
    assert wb.sheetnames == ["SUMMERY", *(ATTRIBUTE_INFO[a].sheet_name for a in ATTRIBUTES)]
    assert [s["name"] for s in cons["workbook"]["sheets"]] == wb.sheetnames[1:]
    for a in ATTRIBUTES:
        info = ATTRIBUTE_INFO[a]
        ws = wb[info.sheet_name]
        assert [c.value for c in ws[1]] == ["S.NO", "TAG NUMBER", "TAG DISCRIPTION", "EQUIPMENT DESCRIPTION",
                                           "SIZE & RATING", f"{info.column_label} IN DATA SHEET (MDS)",
                                           "REMARKS", info.final_label, "STATUS"]
        assert ws.max_row == 718
        for r in range(2, 719):                                        # no comparison output in Phase 1C
            assert ws.cell(row=r, column=7).value is None and ws.cell(row=r, column=8).value is None
            assert ws.cell(row=r, column=9).value is None
    s = wb["SUMMERY"]
    assert [c.value for c in s[1]] == ["SL NO", "DESCRIPTION", "MDS QTY", "REMARKS", "MATCH", "MISMATCH",
                                       "REVIEW REQUIRED"]
    [cov] = cons["coverage"]
    assert [s.cell(row=r, column=3).value for r in range(2, 7)] == [cov["attributes"][a.value]["present"] for a in ATTRIBUTES]
    assert [s.cell(row=r, column=3).value for r in range(2, 7)] == [357, 124, 98, 0, 0]
    assert all(s.cell(row=r, column=c).value is None for r in range(2, 7) for c in range(4, 8))
    assert s["C8"].value == 717


def test_real_mds_columns_equal_original_backup(client, data_dir, real_run):
    """Oracle: the generated MDS column of every sheet equals ORIGINAL BACKUP, except the one documented
    deviation where the backup's human copy dropped punctuation the source file carries."""
    if not ORIGINAL_BACKUP_FILE.exists():
        pytest.skip("ORIGINAL BACKUP not present")
    _, _, _, cons = real_run
    ours, _ = workbook_of(client, data_dir, cons)
    backup = openpyxl.load_workbook(ORIGINAL_BACKUP_FILE, read_only=True, data_only=True)
    seen = 0
    for a in ATTRIBUTES:
        name = ATTRIBUTE_INFO[a].sheet_name
        ref = list(backup[name].iter_rows(values_only=True))
        [col] = [i for i, h in enumerate(ref[0]) if h and "(MDS)" in str(h)]
        mine = list(ours[name].iter_rows(min_row=2, values_only=True))
        assert [r[1] for r in mine] == [r[1] for r in ref[1:] if r[1] is not None]
        for got_row, ref_row in zip(mine, ref[1:]):
            want = None if ref_row[col] in (None, "") else str(ref_row[col])
            if (name, ref_row[1]) in KNOWN_BACKUP_DEVIATIONS:
                assert (got_row[5], want) == KNOWN_BACKUP_DEVIATIONS[(name, ref_row[1])]
                seen += 1
            else:
                assert got_row[5] == want, (name, ref_row[1], got_row[5], want)
    assert seen == len(KNOWN_BACKUP_DEVIATIONS)


def test_real_source_files_unchanged(client, data_dir, real_run, real_datasheet_bytes, real_scope_bytes):
    _, _, src, _ = real_run
    assert REAL_DATASHEET_FILE.read_bytes() == real_datasheet_bytes
    assert REAL_SCOPE_FILE.read_bytes() == real_scope_bytes
    stored = data_dir / "milestones" / "MTL-2026-001" / "sources" / f"{src['sha256']}.xlsx"
    assert hashlib.sha256(stored.read_bytes()).hexdigest() == src["sha256"]


def test_canonical_snapshot_written(client, data_dir, real_run):
    _, _, _, cons = real_run
    [snap] = [o for o in cons["outputs"] if o["kind"] == "CANONICAL_SNAPSHOT"]
    assert snap["stored_path"].startswith("milestones/MTL-2026-001/canonical/run-001/")
    lines = (data_dir / snap["stored_path"]).read_text(encoding="utf-8-sig").splitlines()
    assert len(lines) == 1 + cons["canonical_cell_count"]
    assert any("TOPSAFE CO., LTD" in line for line in lines)


def test_download_serves_the_recorded_artifact(client, real_run):
    _, run, _, cons = real_run
    for art in cons["outputs"]:
        r = client.get(f"/api/runs/{run['id']}/outputs/{art['id']}/download")
        assert r.status_code == 200 and hashlib.sha256(r.content).hexdigest() == art["sha256"]
    assert client.get(f"/api/runs/{run['id']}/outputs/999/download").status_code == 404
    assert [o["id"] for o in client.get(f"/api/runs/{run['id']}/outputs").json()] == [o["id"] for o in cons["outputs"]]


# ================================================================ canonical states (small, synthetic)

def small_run(client, rows: list[list], n: int = 5) -> tuple[dict, dict]:
    m = locked_milestone(client, make_workbook({"S": scope_rows(n)}))
    run = new_run(client, m["id"])
    assert upload(client, f"/api/runs/{run['id']}/sources", ds_file(rows)).status_code == 201
    return m, run


def test_all_canonical_states_distinguished_and_out_of_scope_reported(client):
    rows = [ds_row("TAG-001", as_per_doc="TAG-001", make="ACME"),       # present / empty values
            ds_row("TAG-002"),                                         # row, all values empty
            ds_row("TAG-003", make="A"), ds_row("tag-003", make="B"),  # duplicate
            ds_row("OUTSIDE-9", make="X")]                             # out of scope
    _, run = small_run(client, rows)
    cons = consolidate(client, run["id"])
    assert cons["canonical_row_count"] == 5 and cons["canonical_cell_count"] == 25
    [cov] = cons["coverage"]
    assert cov["out_of_scope_tags"] == [{"tag": "OUTSIDE-9", "rows": [6]}]
    assert cov["duplicate_tags"] == [{"tag": "TAG-003", "rows": 2}]

    rows = {r["tag_number"]: r for r in canonical(client, run["id"])["rows"]}
    assert "OUTSIDE-9" not in rows and len(rows) == 5
    assert rows["TAG-001"]["cells"]["MAKE"]["MDS"]["state"] == "PRESENT"
    assert rows["TAG-001"]["cells"]["MODEL"]["MDS"]["state"] == "ABSENT_VALUE"
    assert rows["TAG-002"]["cells"]["MAKE"]["MDS"]["state"] == "ABSENT_VALUE"
    dup = rows["TAG-003"]["cells"]["MAKE"]["MDS"]
    assert dup["state"] == "CONFLICTING_DUPLICATE" and dup["value"] is None
    assert [p["raw_value"] for p in dup["provenance"]] == ["A", "B"] and [p["row"] for p in dup["provenance"]] == [4, 5]
    assert rows["TAG-003"]["source_rows"] == {"MDS": "DUPLICATE_ROWS"}
    assert rows["TAG-004"]["cells"]["MAKE"]["MDS"] == {"state": "ABSENT_TAG", "value": None, "provenance": []}
    assert rows["TAG-004"]["source_rows"] == {"MDS": "NO_ROW"}
    assert canonical(client, run["id"], filter="duplicates")["filtered_total"] == 1
    assert canonical(client, run["id"], filter="empty_values")["filtered_total"] == 2      # TAG-001, TAG-002
    assert canonical(client, run["id"], filter="no_source_row")["filtered_total"] == 2     # TAG-004, TAG-005


def test_duplicate_is_noted_not_written_in_the_workbook(client, data_dir):
    _, run = small_run(client, [ds_row("TAG-001", make="A"), ds_row("TAG-001", make="B")])
    wb, _ = workbook_of(client, data_dir, consolidate(client, run["id"]))
    c = wb["MAKE COMPARISION"]["F2"]
    assert c.value is None and "none chosen" in c.comment.text and "'A'" in c.comment.text and "'B'" in c.comment.text
    assert wb["SUMMERY"]["C3"].value == 0


def test_consolidation_errors(client):
    m = locked_milestone(client, make_workbook({"S": scope_rows(2)}))
    run = new_run(client, m["id"])
    assert client.post(f"/api/runs/{run['id']}/consolidate").status_code == 409          # nothing uploaded
    upload(client, f"/api/runs/{run['id']}/sources", make_workbook({"X": [["A"], [1]]}))  # FAILED document only
    r = client.post(f"/api/runs/{run['id']}/consolidate")
    assert r.status_code == 409 and "no successfully processed" in r.json()["detail"]
    assert client.get(f"/api/runs/{run['id']}/consolidation").status_code == 404
    assert client.get(f"/api/runs/{run['id']}/canonical").status_code == 404
    assert client.post("/api/runs/999/consolidate").status_code == 404
    upload(client, f"/api/runs/{run['id']}/sources", ds_file([ds_row("TAG-001")]))
    cons = consolidate(client, run["id"])
    assert [d["included"] for d in cons["documents"]] == [False, True]
    assert cons["documents"][0]["excluded_reason"] == "processing failed"
    assert client.get(f"/api/runs/{run['id']}/canonical?filter=bogus").status_code == 422


def test_consolidation_alone_persists_no_comparison_state(client):
    _, run = small_run(client, [ds_row("TAG-001", make="ACME")])
    consolidate(client, run["id"])
    con = db(client)
    assert con.execute("select count(*) from comparison_result").fetchone()[0] == 0
    assert con.execute("select count(*) from run_comparison").fetchone()[0] == 0
    cols = {r[1] for r in con.execute("pragma table_info(canonical_cell)")}
    assert not cols & {"status", "final_value", "remarks"}
    assert {r[0] for r in con.execute("select distinct state from canonical_cell")} <= \
        {"PRESENT", "ABSENT_VALUE", "ABSENT_TAG", "CONFLICTING_DUPLICATE"}


# ================================================================ Test B — two sources in one run

def test_two_sources_in_one_run(client, data_dir, spir_fixture):
    m = locked_milestone(client, make_workbook({"S": scope_rows(4)}))
    run = new_run(client, m["id"])
    ds = upload(client, f"/api/runs/{run['id']}/sources",
                ds_file([ds_row("TAG-001", as_per_doc="TAG-001", make="ACME"), ds_row("TAG-002", make="MDS ONLY")])).json()
    sp = upload(client, f"/api/runs/{run['id']}/sources",
                fixture_workbook([fixture_row("TAG-001", as_per_doc="T-001", make="ACME CORP", serial="SN-1"),
                                  fixture_row("TAG-003", make="SPIR ONLY")]), name="fixture.xlsx").json()
    assert (ds["source_type"], sp["source_type"]) == ("MDS", "SPIR")
    assert ds["run_id"] == sp["run_id"] == run["id"] and ds["document_id"] != sp["document_id"]

    cons = consolidate(client, run["id"])
    assert [t["code"] for t in cons["source_types"]] == ["MDS", "SPIR"]
    assert cons["canonical_row_count"] == 4 and cons["canonical_cell_count"] == 4 * 5 * 2
    assert {c["source_type"]: c["rows"] for c in cons["coverage"]} == {"MDS": 2, "SPIR": 2}

    rows = {r["tag_number"]: r for r in canonical(client, run["id"])["rows"]}
    make = rows["TAG-001"]["cells"]["MAKE"]
    assert (make["MDS"]["value"], make["SPIR"]["value"]) == ("ACME", "ACME CORP")      # neither overwritten
    assert make["MDS"]["provenance"][0]["document"] != make["SPIR"]["provenance"][0]["document"]
    assert rows["TAG-002"]["source_rows"] == {"MDS": "ROW", "SPIR": "NO_ROW"}
    assert rows["TAG-003"]["source_rows"] == {"MDS": "NO_ROW", "SPIR": "ROW"}
    assert rows["TAG-004"]["source_rows"] == {"MDS": "NO_ROW", "SPIR": "NO_ROW"}

    con = db(client)
    per_doc = dict(con.execute("select c.source_type, count(distinct r.run_source_id) from canonical_cell c "
                               "join source_record r on r.id = c.source_record_id group by c.source_type"))
    assert per_doc == {"MDS": 1, "SPIR": 1}                          # records stay separately identifiable

    wb, _ = workbook_of(client, data_dir, cons)
    ws = wb["MAKE COMPARISION"]
    assert [c.value for c in ws[1]][5:] == ["MAKE IN DATA SHEET (MDS)", "MAKE IN SPIR (MIR)", "REMARKS", "FINAL MAKE", "STATUS"]
    assert [ws.cell(row=2, column=c).value for c in range(6, 11)] == ["ACME", "ACME CORP", None, None, None]
    assert [c.value for c in wb["SUMMERY"][1]][2:4] == ["MDS QTY", "MIR QTY"]
    assert [c.value for c in wb["SUMMERY"][3]][2:4] == [2, 2]
    headers = [c.value for ws in wb.worksheets for c in ws[1]]
    assert not any(h and ("GA" in h.split() or "MXS" in h or "MTC" in h or "MOM" in h or "ASSET" in h) for h in headers)


# ================================================================ Test C — run isolation

def _run_state(con, run_id: int) -> tuple:
    return (con.execute("select id, milestone_tag_id, attribute, source_type, state, raw_value, source_record_ids "
                        "from canonical_cell where run_id = ? order by id", (run_id,)).fetchall(),
            con.execute("select id, stored_path, sha256 from output_artifact where run_id = ? order by id", (run_id,)).fetchall(),
            con.execute("select count(*) from source_record where run_id = ?", (run_id,)).fetchone(),
            con.execute("select status from processing_run where id = ?", (run_id,)).fetchone())


def test_earlier_run_unchanged_when_later_run_adds_a_source(client, data_dir, spir_fixture):
    m = locked_milestone(client, make_workbook({"S": scope_rows(4)}))
    run1 = new_run(client, m["id"])
    data = ds_file([ds_row("TAG-001", make="ACME"), ds_row("TAG-002", make="MDS-2")])
    upload(client, f"/api/runs/{run1['id']}/sources", data, name="ds.xlsx")
    cons1 = consolidate(client, run1["id"])
    con = db(client)
    before = _run_state(con, run1["id"])
    files_before = {o["stored_path"]: (data_dir / o["stored_path"]).read_bytes() for o in cons1["outputs"]}

    # A consolidated run is frozen.
    r = upload(client, f"/api/runs/{run1['id']}/sources", fixture_workbook([fixture_row("TAG-001")]))
    assert r.status_code == 409 and "new run" in r.json()["detail"]
    assert client.post(f"/api/runs/{run1['id']}/consolidate").status_code == 409

    # Run 2 starts from run 1's documents and adds a second source.
    r = client.post(f"/api/milestones/{m['id']}/runs", json={"base_run_id": run1["id"], "notes": "add SPIR"})
    assert r.status_code == 201, r.text
    run2 = r.json()
    assert [s["source_type"] for s in run2["sources"]] == ["MDS"]
    assert run2["sources"][0]["document_id"] == cons1["documents"][0]["document_id"]      # same immutable document
    upload(client, f"/api/runs/{run2['id']}/sources",
           fixture_workbook([fixture_row("TAG-001", make="ACME SPIR"), fixture_row("TAG-004", make="S-4")]))
    cons2 = consolidate(client, run2["id"])
    assert [t["code"] for t in cons2["source_types"]] == ["MDS", "SPIR"]
    assert cons2["canonical_cell_count"] == 2 * cons1["canonical_cell_count"]
    assert all("/run-002/" in o["stored_path"] for o in cons2["outputs"])

    # Run 1: database rows, files and status exactly as before.
    assert _run_state(con, run1["id"]) == before
    for path, content in files_before.items():
        assert (data_dir / path).read_bytes() == content
    assert client.get(f"/api/runs/{run1['id']}/consolidation").json()["source_types"] == cons1["source_types"]
    assert con.execute("select count(*) from source_document where sha256 = ?",
                       (hashlib.sha256(data).hexdigest(),)).fetchone()[0] == 1


def test_base_run_must_belong_to_the_milestone(client):
    a = locked_milestone(client, make_workbook({"S": scope_rows(2)}), code="A")
    b = locked_milestone(client, make_workbook({"S": scope_rows(2)}), code="B")
    run_a = new_run(client, a["id"])
    r = client.post(f"/api/milestones/{b['id']}/runs", json={"base_run_id": run_a["id"]})
    assert r.status_code == 422


def test_removing_a_source_opens_a_new_run_without_it(client, spir_fixture):
    m = locked_milestone(client, make_workbook({"S": scope_rows(4)}))
    run1 = new_run(client, m["id"])
    upload(client, f"/api/runs/{run1['id']}/sources", ds_file([ds_row("TAG-001", make="ACME")]), name="ds.xlsx")
    upload(client, f"/api/runs/{run1['id']}/sources", fixture_workbook([fixture_row("TAG-002", make="S-2")]))
    cons1 = consolidate(client, run1["id"])
    con = db(client)
    before = _run_state(con, run1["id"])
    mds, spir = client.get(f"/api/runs/{run1['id']}").json()["sources"]

    r = client.delete(f"/api/runs/{run1['id']}/sources/{spir['id']}")
    assert r.status_code == 200, r.text
    run2 = r.json()
    assert run2["id"] != run1["id"] and run2["status"] != "CONSOLIDATED"
    assert [(s["source_type"], s["document_id"]) for s in run2["sources"]] == [("MDS", mds["document_id"])]
    assert client.get(f"/api/milestones/{m['id']}/runs").json()[0]["id"] == run2["id"]    # now the current run

    # The removed-from run, its evidence and its outputs are untouched.
    assert _run_state(con, run1["id"]) == before
    assert client.get(f"/api/runs/{run1['id']}/consolidation").json()["source_types"] == cons1["source_types"]

    assert client.delete(f"/api/runs/{run2['id']}/sources/{spir['id']}").status_code == 404   # not part of run 2
    assert client.delete(f"/api/runs/999999/sources/{mds['id']}").status_code == 404


def test_consolidated_evidence_is_immutable_in_the_database(client):
    _, run = small_run(client, [ds_row("TAG-001", make="ACME")])
    consolidate(client, run["id"])
    con = db(client)
    doc_id, = con.execute("select id from source_document").fetchone()
    for sql in ("UPDATE canonical_cell SET raw_value = 'X'", "DELETE FROM canonical_cell",
                "UPDATE run_consolidation SET cell_count = 0", "DELETE FROM output_artifact",
                "UPDATE output_artifact SET sha256 = 'x'"):
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            con.execute(sql)
    with pytest.raises(sqlite3.IntegrityError, match="frozen"):
        con.execute("INSERT INTO run_source (run_id, source_document_id, processing_status, report, processed_at) "
                    "VALUES (?, ?, 'PROCESSED', '{}', '2026-01-01')", (run["id"], doc_id))
    with pytest.raises(sqlite3.IntegrityError, match="earlier state"):
        con.execute("UPDATE processing_run SET status = 'DOCUMENTS_UPLOADED' WHERE id = ?", (run["id"],))


def test_failed_consolidation_leaves_nothing_behind_and_can_be_retried(client, data_dir, monkeypatch):
    _, run = small_run(client, [ds_row("TAG-001", make="ACME")])
    files = client.app.state.consolidation_service.files
    real_store = files.store_generated
    calls = []

    def failing_store(*args):
        calls.append(args)
        if len(calls) == 2:                                   # the workbook write fails after the snapshot
            raise OSError("disk full")
        return real_store(*args)

    monkeypatch.setattr(files, "store_generated", failing_store)
    with pytest.raises(OSError):
        client.post(f"/api/runs/{run['id']}/consolidate")
    monkeypatch.setattr(files, "store_generated", real_store)
    con = db(client)
    assert con.execute("select count(*) from canonical_cell").fetchone()[0] == 0
    assert con.execute("select count(*) from output_artifact").fetchone()[0] == 0
    assert client.get(f"/api/runs/{run['id']}").json()["status"] == "DOCUMENTS_UPLOADED"
    assert not any((data_dir / "milestones" / "MTL-2026-001" / "canonical").rglob("*.csv"))
    cons = consolidate(client, run["id"])                     # retry succeeds
    assert {o["stored_path"].rsplit("/", 1)[1] for o in cons["outputs"]} == {"canonical.csv", "consolidated.xlsx"}


def test_consolidation_record_json_is_complete(client):
    _, run = small_run(client, [ds_row("TAG-001", make="ACME")])
    consolidate(client, run["id"])
    types, ids = db(client).execute("select source_types, run_source_ids from run_consolidation").fetchone()
    assert json.loads(types) == ["MDS"] and len(json.loads(ids)) == 1


# ================================================================ Test D — real Data Sheet + real SPIR

def row_of(client, run_id: int, tag: str) -> dict:
    [row] = [r for r in canonical(client, run_id, search=tag)["rows"] if r["tag_number"] == tag]
    return row


def _ds_keys(data: bytes) -> set[str]:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = list(wb["DATASHEET"].iter_rows(min_row=2, values_only=True))
    wb.close()
    return {str(r[1]).strip().upper() for r in rows if r[1] is not None and str(r[1]).strip()}


@pytest.fixture
def real_two_source_run(client, real_scope_bytes, real_datasheet_bytes, real_spir_bytes):
    m = locked_milestone(client, real_scope_bytes)
    run = new_run(client, m["id"])
    ds = upload(client, f"/api/runs/{run['id']}/sources", real_datasheet_bytes, name=REAL_DATASHEET_FILE.name).json()
    sp = upload(client, f"/api/runs/{run['id']}/sources", real_spir_bytes, name=REAL_SPIR_FILE.name).json()
    return m, run, ds, sp, consolidate(client, run["id"])


def test_real_spir_upload_is_recognised_and_processed(real_two_source_run):
    _, _, ds, sp, _ = real_two_source_run
    assert (ds["source_type"], sp["source_type"]) == ("MDS", "SPIR")          # auto-identified, no type given
    assert sp["processing_status"] == "PROCESSED" and sp["errors"] == [] and sp["warnings"] == []
    assert (sp["sheet_name"], sp["row_count"], sp["unique_tag_count"], sp["matched_tag_count"],
            sp["unmatched_tag_count"], sp["duplicate_tag_count"]) == ("SPIR DATA", 464, 464, 464, 0, 0)
    assert sp["attributes_provided"] == [a.value for a in ATTRIBUTES]
    assert sp["attribute_counts"] == {"TAG_NUMBER": 377, "MAKE": 458, "MODEL": 355, "SERIAL_NUMBER": 331,
                                      "PART_NUMBER": 370}


def test_real_datasheet_and_spir_consolidate_on_the_717_tag_scope(client, real_two_source_run, real_scope_bytes,
                                                                     real_datasheet_bytes, real_spir_bytes):
    m, run, ds, sp, cons = real_two_source_run
    scope = scope_tags_of(real_scope_bytes)
    n = len(scope)
    assert cons["scope_tag_count"] == cons["canonical_row_count"] == m["tag_count"] == n == 717
    assert [t["code"] for t in cons["source_types"]] == ["MDS", "SPIR"]
    assert cons["canonical_cell_count"] == n * len(ATTRIBUTES) * 2

    cov = {c["source_type"]: c for c in cons["coverage"]}
    for code, src in (("MDS", ds), ("SPIR", sp)):
        c = cov[code]
        assert c["rows"] == src["row_count"] and c["in_scope_tags"] == src["matched_tag_count"]
        assert c["scope_tags_without_row"] == n - c["in_scope_tags"]
        assert c["attributes_provided"] == [a.value for a in ATTRIBUTES]
        for a in ATTRIBUTES:
            st = c["attributes"][a.value]
            assert st["present"] == src["attribute_counts"][a.value]
            assert st["present"] + st["absent_value"] + st["absent_tag"] + st["duplicate"] == n
    # MDS golden figures are unchanged by adding a second source.
    assert (cov["MDS"]["rows"], cov["MDS"]["in_scope_tags"], cov["MDS"]["scope_tags_without_row"]) == (436, 436, 281)
    assert {a: v["present"] for a, v in cov["MDS"]["attributes"].items()} == \
        {"TAG_NUMBER": 357, "MAKE": 124, "MODEL": 98, "SERIAL_NUMBER": 0, "PART_NUMBER": 0}
    # SPIR regression figures.
    assert (cov["SPIR"]["rows"], cov["SPIR"]["in_scope_tags"], cov["SPIR"]["out_of_scope_tag_count"],
            cov["SPIR"]["scope_tags_without_row"]) == (464, 464, 0, 253)
    assert {a: v["present"] for a, v in cov["SPIR"]["attributes"].items()} == \
        {"TAG_NUMBER": 377, "MAKE": 458, "MODEL": 355, "SERIAL_NUMBER": 331, "PART_NUMBER": 370}
    assert (cov["SPIR"]["tag_as_per_doc_equal"], cov["SPIR"]["tag_as_per_doc_differs"]) == (357, 20)

    # Canonical rows are the scope, in scope order; source coverage per row is derived from both files.
    page = canonical(client, run["id"])
    assert [r["tag_number"] for r in page["rows"]] == scope
    spir_keys = {str(e["TAG NUMBER"]).strip().upper() for e in independent_spir(real_spir_bytes)}
    ds_keys = _ds_keys(real_datasheet_bytes)
    neither = [t for t in scope if t.upper() not in ds_keys | spir_keys]
    assert canonical(client, run["id"], filter="no_source_row")["filtered_total"] == len(neither)
    only_one = [t for t in scope if (t.upper() in ds_keys) != (t.upper() in spir_keys)]
    assert canonical(client, run["id"], filter="missing_in_some_source")["filtered_total"] == len(only_one)
    for r in page["rows"]:
        k = r["tag_number"].upper()
        assert r["source_rows"] == {"MDS": "ROW" if k in ds_keys else "NO_ROW",
                                    "SPIR": "ROW" if k in spir_keys else "NO_ROW"}


def test_real_two_source_values_keep_source_and_provenance(client, real_two_source_run):
    _, run, _, _, _ = real_two_source_run
    row = row_of(client, run["id"], "88-GV-1601")
    make = row["cells"]["MAKE"]["SPIR"]
    assert (make["state"], make["value"]) == ("PRESENT", "JIANGSU YDF VALVE CO.,LTD")      # source value kept
    [p] = make["provenance"]
    assert (p["document"], p["sheet"]) == (REAL_SPIR_FILE.name, "SPIR DATA") and p["cell"].startswith("N")
    assert p["document_reference"] and p["page"]
    row = row_of(client, run["id"], "68-FIT-1987")
    assert row["cells"]["TAG_NUMBER"]["SPIR"]["value"] == "68-FT-1987"                      # milestone tag kept
    assert row["tag_number"] == "68-FIT-1987"


def test_real_two_source_workbook_columns_equal_original_backup(client, data_dir, real_two_source_run):
    """Oracle: the generated MDS and SPIR columns of every sheet equal ORIGINAL BACKUP, except the two
    documented deviations; no FINAL / STATUS is written (no comparison in this phase)."""
    if not ORIGINAL_BACKUP_FILE.exists():
        pytest.skip("ORIGINAL BACKUP not present")
    _, _, _, _, cons = real_two_source_run
    ours, _ = workbook_of(client, data_dir, cons)
    backup = openpyxl.load_workbook(ORIGINAL_BACKUP_FILE, read_only=True, data_only=True)
    deviations = KNOWN_BACKUP_DEVIATIONS | KNOWN_SPIR_BACKUP_DEVIATIONS
    seen = 0
    for a in ATTRIBUTES:
        info = ATTRIBUTE_INFO[a]
        ws = ours[info.sheet_name]
        assert [c.value for c in ws[1]] == ["S.NO", "TAG NUMBER", "TAG DISCRIPTION", "EQUIPMENT DESCRIPTION",
                                           "SIZE & RATING", f"{info.column_label} IN DATA SHEET (MDS)",
                                           f"{info.column_label} IN SPIR (MIR)", "REMARKS", info.final_label, "STATUS"]
        ref = list(backup[info.sheet_name].iter_rows(values_only=True))
        [mds_col] = [i for i, h in enumerate(ref[0]) if h and "(MDS)" in str(h)]
        [spir_col] = [i for i, h in enumerate(ref[0]) if h and "SPIR" in str(h)]
        mine = list(ws.iter_rows(min_row=2, values_only=True))
        assert [r[1] for r in mine] == [r[1] for r in ref[1:] if r[1] is not None]
        for got, want_row in zip(mine, ref[1:]):
            for got_v, col in ((got[5], mds_col), (got[6], spir_col)):
                want = None if want_row[col] in (None, "") else str(want_row[col])
                if got_v != want:
                    assert deviations[(info.sheet_name, want_row[1])] == (got_v, want)
                    seen += 1
            assert got[7] is None and got[8] is None and got[9] is None               # REMARKS / FINAL / STATUS
    assert seen == len(deviations)
    s = ours["SUMMERY"]
    assert [c.value for c in s[1]][2:4] == ["MDS QTY", "MIR QTY"]
    assert [s.cell(row=r, column=4).value for r in range(2, 7)] == [377, 458, 355, 331, 370]


def test_source_missing_a_column_consolidates_without_it(client, data_dir, real_scope_bytes, real_datasheet_bytes,
                                                          real_spir_bytes):
    m = locked_milestone(client, real_scope_bytes)
    run = new_run(client, m["id"])
    upload(client, f"/api/runs/{run['id']}/sources", real_datasheet_bytes, name=REAL_DATASHEET_FILE.name)
    data = without_column(real_spir_bytes, "SPIR DATA", "PART NUMBER")
    sp = upload(client, f"/api/runs/{run['id']}/sources", data, name="spir-no-part.xlsx").json()
    assert sp["source_type"] == "SPIR" and sp["processing_status"] == "PROCESSED_WITH_WARNINGS"
    assert [w["code"] for w in sp["warnings"]] == ["FIELD_NOT_PROVIDED"]
    assert "PART_NUMBER" not in sp["attributes_provided"] and "PART_NUMBER" not in sp["attribute_counts"]

    cons = consolidate(client, run["id"])
    n = cons["canonical_row_count"]
    assert n == m["tag_count"]
    assert cons["canonical_cell_count"] == n * len(ATTRIBUTES) + n * (len(ATTRIBUTES) - 1)
    spir = next(c for c in cons["coverage"] if c["source_type"] == "SPIR")
    assert "PART_NUMBER" not in spir["attributes_provided"] and "PART_NUMBER" not in spir["attributes"]
    row = row_of(client, run["id"], "6834-P-64A")
    assert set(row["cells"]["PART_NUMBER"]) == {"MDS"}                  # no SPIR cell, not ABSENT_TAG
    assert row["cells"]["SERIAL_NUMBER"]["SPIR"]["value"] == "665436"
    assert row["source_rows"] == {"MDS": row["source_rows"]["MDS"], "SPIR": "ROW"}

    wb, _ = workbook_of(client, data_dir, cons)
    ws = wb["PART NUMBER COMPARISION"]
    assert [c.value for c in ws[1]][5:7] == ["P/N IN DATA SHEET (MDS)", "P/N IN SPIR (MIR)"]
    assert all(ws.cell(row=r, column=7).value is None for r in range(2, n + 2))   # left empty, nothing invented


def test_out_of_scope_spir_tag_never_expands_the_scope(client, real_scope_bytes, real_spir_bytes):
    m = locked_milestone(client, real_scope_bytes)
    run = new_run(client, m["id"])
    data = with_extra_row(real_spir_bytes, "SPIR DATA", {"TAG NUMBER": "NOT-IN-MILESTONE-999", "MAKE": "X"})
    sp = upload(client, f"/api/runs/{run['id']}/sources", data, name="spir-plus-one.xlsx").json()
    assert (sp["matched_tag_count"], sp["unmatched_tag_count"]) == (464, 1)
    assert "OUT_OF_SCOPE_TAG" in [w["code"] for w in sp["warnings"]]
    cons = consolidate(client, run["id"])
    assert cons["canonical_row_count"] == m["tag_count"]
    [cov] = cons["coverage"]
    assert [t["tag"] for t in cov["out_of_scope_tags"]] == ["NOT-IN-MILESTONE-999"]
    page = canonical(client, run["id"])
    assert page["total"] == m["tag_count"]
    assert canonical(client, run["id"], search="NOT-IN-MILESTONE")["filtered_total"] == 0
