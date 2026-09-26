"""Data Sheet adapter: identification, semantic field discovery, extraction, and the real-file oracle.

Column positions of the real file appear ONLY here, as fixture assertions — never in the adapter.
"""

import io
from collections import Counter

import openpyxl
import pytest

from backend.domain.sources import Attribute
from backend.extraction import registry
from backend.extraction.datasheet import DATASHEET
from backend.extraction.tabular import load_workbook
from tests.conftest import DS_HEADERS, ORIGINAL_BACKUP_FILE, REAL_DATASHEET_FILE, ds_row, make_workbook


def run(data: bytes, name: str = "ds.xlsx", source_type: str | None = None):
    book = load_workbook(name, data)
    ident = registry.identify(book, source_type)
    return ident, (registry.extract(book, ident) if ident.source_type else None)


def codes(issues) -> list[str]:
    return [i.code for i in issues]


def by_tag(ext) -> dict[str, dict[Attribute, str | None]]:
    return {r.normalized_key: {a: v.raw for a, v in r.values.items()} for r in ext.rows}


# ---------------------------------------------------------------- real file

def independent_datasheet(data: bytes) -> list[dict]:
    """Read DATASHEET with plain openpyxl and stripped-header lookup — a separate code path."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = list(wb["DATASHEET"].iter_rows(values_only=True))
    wb.close()
    idx = {str(h).strip(): i for i, h in enumerate(rows[0]) if h is not None}
    out = []
    for r in rows[1:]:
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in r):
            continue
        out.append({k: r[idx[k]] for k in ("TAG NUMBER", "TAG NUMBER AS PER DOC", "MAKE", "MODEL",
                                            "SERIAL NUMBER", "PART NUMBER")})
    return out


def filled(v) -> bool:
    return v is not None and str(v).strip() != ""


def test_real_datasheet_identified_and_other_sheets_reported(real_datasheet_bytes):
    ident, ext = run(real_datasheet_bytes, REAL_DATASHEET_FILE.name)
    assert (ident.source_type, ident.sheet_name) == ("MDS", "DATASHEET")
    assert {s["sheet"] for s in ident.skipped_sheets} == {"PRD COUNT", "Sheet1"}
    assert ext.ok and ext.issues == []


def test_real_datasheet_mapping_fixture_positions(real_datasheet_bytes):
    """Positions verified in the real file. Asserted here; never encoded in the adapter."""
    _, ext = run(real_datasheet_bytes, REAL_DATASHEET_FILE.name)
    cols = {b.key: b.column for b in ext.mapping}
    assert cols == {"TAG_KEY": "B", "TAG_AS_PER_DOC": "G", "MAKE": "J", "MODEL": "K", "SERIAL_NUMBER": "L",
                    "PART_NUMBER": "M", "DOCUMENT_REFERENCE": "F", "DOCUMENT_IDB": "E", "PAGE_NO": "H"}
    headers = {b.key: b.header for b in ext.mapping}
    assert headers["TAG_AS_PER_DOC"] == " TAG NUMBER AS PER DOC"     # leading space, matched semantically


def test_real_datasheet_counts_derived_from_file(real_datasheet_bytes):
    expected = independent_datasheet(real_datasheet_bytes)
    _, ext = run(real_datasheet_bytes, REAL_DATASHEET_FILE.name)
    assert len(ext.rows) == len(expected)
    assert len({r.normalized_key for r in ext.rows}) == len({str(e["TAG NUMBER"]).strip().upper() for e in expected})
    counts = Counter(a for r in ext.rows for a, v in r.values.items() if v.normalized is not None)
    for attr, header in [(Attribute.TAG_NUMBER, "TAG NUMBER AS PER DOC"), (Attribute.MAKE, "MAKE"),
                         (Attribute.MODEL, "MODEL"), (Attribute.SERIAL_NUMBER, "SERIAL NUMBER"),
                         (Attribute.PART_NUMBER, "PART NUMBER")]:
        assert counts.get(attr, 0) == sum(1 for e in expected if filled(e[header])), attr


def test_real_datasheet_regression_figures(real_datasheet_bytes):
    _, ext = run(real_datasheet_bytes, REAL_DATASHEET_FILE.name)
    counts = Counter(a for r in ext.rows for a, v in r.values.items() if v.normalized is not None)
    assert len(ext.rows) == 436 and len({r.normalized_key for r in ext.rows}) == 436
    assert [counts.get(a, 0) for a in Attribute] == [357, 124, 98, 0, 0]
    tag_vals = [(r.normalized_key, r.values[Attribute.TAG_NUMBER].normalized) for r in ext.rows
                if r.values[Attribute.TAG_NUMBER].normalized]
    assert sum(1 for k, v in tag_vals if v.upper() == k) == 343
    assert sum(1 for k, v in tag_vals if v.upper() != k) == 14


def test_join_key_and_tag_as_per_doc_are_not_conflated(real_datasheet_bytes):
    _, ext = run(real_datasheet_bytes, REAL_DATASHEET_FILE.name)
    rows = {r.normalized_key: r for r in ext.rows}
    assert rows["68-SP-051"].values[Attribute.TAG_NUMBER].normalized == "SP-123"
    assert rows["88-LT-802"].values[Attribute.TAG_NUMBER].normalized == "88-LIT-802"


def test_raw_values_preserved_and_only_trimmed_when_normalized(real_datasheet_bytes):
    _, ext = run(real_datasheet_bytes, REAL_DATASHEET_FILE.name)
    rows = {r.row_number: r for r in ext.rows}
    make = rows[256].values[Attribute.MAKE]
    assert (make.raw, make.normalized, make.column) == (" Honeywell", "Honeywell", "J")
    model = rows[391].values[Attribute.MODEL]
    assert (model.raw, model.normalized, model.raw_type) == ("200", "200", "int")
    assert rows[2].provenance == {"DATASHEET REFERENCE-ORIGINAL": "4391-MTY-5-13-0101-D",
                                  "DATASHEET REFERENCE-DOC IDB": "NG-MDS-VABP-400021", "PAGE NO": "7"}
    assert {k: v for k, v in ext.location_of(rows[2]).items() if k != "source_fields"} == \
        {"document_reference": "4391-MTY-5-13-0101-D", "document_idb": "NG-MDS-VABP-400021", "page": "7"}


# The one cell where ORIGINAL BACKUP does not reproduce the Data Sheet verbatim. The source file (cell J428)
# reads 'TOPSAFE CO., LTD'; the human backup has 'TOPSAFE CO LTD'. The adapter must keep the source value —
# punctuation stripping is a judgement the spec forbids — so this is pinned as a known oracle deviation.
# See docs/SOURCE_ADAPTERS.md. Any *other* difference fails the test.
KNOWN_BACKUP_DEVIATIONS = {("MAKE COMPARISION", "68-SP-015"): ("TOPSAFE CO., LTD", "TOPSAFE CO LTD")}


def test_extraction_reproduces_original_backup_mds_columns(real_datasheet_bytes):
    """Oracle: every MDS cell of ORIGINAL BACKUP equals the extracted raw value for that tag (or both empty),
    except the documented deviation above."""
    if not ORIGINAL_BACKUP_FILE.exists():
        pytest.skip("ORIGINAL BACKUP not present")
    _, ext = run(real_datasheet_bytes, REAL_DATASHEET_FILE.name)
    extracted = by_tag(ext)
    sheets = {"TAG NUMBER COMPARISION": Attribute.TAG_NUMBER, "MAKE COMPARISION": Attribute.MAKE,
              "MODEL COMPARISION": Attribute.MODEL, "SERIAL NUMBER COMPARISION": Attribute.SERIAL_NUMBER,
              "PART NUMBER COMPARISION": Attribute.PART_NUMBER}
    wb = openpyxl.load_workbook(ORIGINAL_BACKUP_FILE, read_only=True, data_only=True)
    deviations_seen = 0
    try:
        for sheet, attr in sheets.items():
            rows = list(wb[sheet].iter_rows(values_only=True))
            [mds_col] = [i for i, h in enumerate(rows[0]) if h and "(MDS)" in str(h)]
            compared = 0
            for r in rows[1:]:
                if r[1] is None:
                    continue
                oracle = r[mds_col]
                oracle = None if oracle is None or str(oracle) == "" else str(oracle)
                ours = extracted.get(str(r[1]).strip().upper(), {}).get(attr)
                if (sheet, r[1]) in KNOWN_BACKUP_DEVIATIONS:
                    assert (ours, oracle) == KNOWN_BACKUP_DEVIATIONS[(sheet, r[1])]
                    deviations_seen += 1
                else:
                    assert ours == oracle, (sheet, r[1], ours, oracle)
                compared += 1
            assert compared == 717
        assert deviations_seen == len(KNOWN_BACKUP_DEVIATIONS)
    finally:
        wb.close()


def test_reordered_real_columns_extract_identically(real_datasheet_bytes):
    """Reverse every column of the real DATASHEET: same values, different letters."""
    wb = openpyxl.load_workbook(io.BytesIO(real_datasheet_bytes), data_only=True)
    ws = wb["DATASHEET"]
    grid = [[c.value for c in row] for row in ws.iter_rows()]
    wb.remove(ws)
    new = wb.create_sheet("DATASHEET", 0)
    for row in grid:
        new.append(list(reversed(row)))
    buf = io.BytesIO()
    wb.save(buf)

    _, original = run(real_datasheet_bytes, REAL_DATASHEET_FILE.name)
    _, reordered = run(buf.getvalue(), "reordered.xlsx")
    assert reordered.ok
    assert by_tag(reordered) == by_tag(original)
    orig_cols = {b.key: b.column for b in original.mapping}
    new_cols = {b.key: b.column for b in reordered.mapping}
    assert all(orig_cols[k] != new_cols[k] for k in ("TAG_KEY", "MAKE", "MODEL"))


# ---------------------------------------------------------------- synthetic layouts

def ds_book(headers: list, rows: list[list], sheet: str = "DATASHEET", extra_sheets: dict | None = None) -> bytes:
    return make_workbook({sheet: [headers] + rows, **(extra_sheets or {})})


def test_harmless_extra_and_missing_irrelevant_columns():
    headers = ["TAG NUMBER", "NOTES", " TAG NUMBER AS PER DOC", "MAKE", "WEIGHT", "MODEL", "SERIAL NUMBER",
               "PART NUMBER", "VALVE SCHEDULE SIZE & RATING"]
    _, ext = run(ds_book(headers, [["T-1", "n", "T-1", "ACME", 5, "M1", None, None, "x"]]))
    assert ext.ok
    assert by_tag(ext)["T-1"] == {Attribute.TAG_NUMBER: "T-1", Attribute.MAKE: "ACME", Attribute.MODEL: "M1",
                                  Attribute.SERIAL_NUMBER: None, Attribute.PART_NUMBER: None}


def test_documented_aliases_are_accepted():
    headers = ["TAG NO.", "Tag Number As Per Document", "MANUFACTURER", "MODEL NO", "SERIAL NO", "PART NO"]
    _, ext = run(ds_book(headers, [["t-1", "T-1", "ACME", "M1", "S1", "P1"]]))
    assert ext.ok and by_tag(ext)["T-1"][Attribute.PART_NUMBER] == "P1"


def test_missing_attribute_column_is_reported_not_guessed_and_does_not_fail():
    """Available-data principle (Phase 1D-1): a document without a MAKE column still processes; MAKE is
    reported as not provided and nothing is filled in. (Before 1D-1 every attribute column was required.)"""
    headers = ["TAG NUMBER", " TAG NUMBER AS PER DOC", "MODEL", "SERIAL NUMBER", "PART NUMBER", "REMARKS"]
    _, ext = run(ds_book(headers, [["T-1", "T-1", "M", None, None, "ok"]]))
    assert ext.ok and len(ext.rows) == 1
    [warn] = ext.issues
    assert warn.code == "FIELD_NOT_PROVIDED" and warn.value == "MAKE"
    assert "DATASHEET" in warn.message and "'REMARKS'" in warn.message     # sheet and headers found
    assert Attribute.MAKE not in ext.provided_attributes
    assert Attribute.MAKE not in ext.rows[0].values                        # absent, not an empty placeholder
    assert ext.rows[0].values[Attribute.MODEL].raw == "M"


def test_document_with_no_attribute_column_fails():
    _, ext = run(ds_book(["TAG NUMBER", "REMARKS"], [["T-1", "x"]]))
    assert not ext.ok and ext.rows == [] and codes(ext.issues) == ["NO_ATTRIBUTE_FIELD"]


def test_ambiguous_field_mapping_is_refused():
    headers = ["TAG NUMBER", " TAG NUMBER AS PER DOC", "MAKE", "MANUFACTURER", "MODEL", "SERIAL NUMBER", "PART NUMBER"]
    _, ext = run(ds_book(headers, [["T-1", "T-1", "A", "B", "M", None, None]]), name="amb.xlsx")
    assert not ext.ok and ext.rows == []
    [err] = ext.issues
    assert err.code == "AMBIGUOUS_FIELD" and err.value == "MAKE"
    assert "C 'MAKE'" in err.message and "D 'MANUFACTURER'" in err.message and "amb.xlsx" in err.message


def test_substring_headers_do_not_bind():
    headers = ["TAG NUMBER", " TAG NUMBER AS PER DOC", "MAKE", "MODEL", "SERIAL NUMBER", "PART NUMBER REMARKS"]
    _, ext = run(ds_book(headers, [["T-1", None, None, None, None, "x"]]))
    assert codes(ext.issues) == ["FIELD_NOT_PROVIDED"] and ext.issues[0].value == "PART NUMBER"
    assert Attribute.PART_NUMBER not in ext.provided_attributes


def test_tag_as_per_doc_never_used_as_join_key():
    headers = [" TAG NUMBER AS PER DOC", "MAKE", "MODEL", "SERIAL NUMBER", "PART NUMBER"]
    _, ext = run(ds_book(headers, [["T-1", "A", None, None, None]]))
    assert codes(ext.issues) == ["MISSING_FIELD"] and "join key" in ext.issues[0].message


def test_duplicate_keys_are_all_kept():
    rows = [ds_row("T-1", make="A"), ds_row("T-1", make="B"), ds_row("t-1 ", make="C")]
    _, ext = run(ds_book(DS_HEADERS, rows))
    assert [r.values[Attribute.MAKE].raw for r in ext.rows if r.normalized_key == "T-1"] == ["A", "B", "C"]


def test_workbook_naming_no_document_type_is_not_identified():
    ident, ext = run(make_workbook({"Export": [["TAG NUMBER", "MAKE"], ["T-1", "A"]]}))
    assert ext is None and codes(ident.issues) == ["SOURCE_TYPE_UNKNOWN"]
    ident, ext = run(make_workbook({"SPIR": [["TAG NUMBER", "MAKE"], ["T-1", "A"]]}))       # named by keyword
    assert ident.source_type == "SPIR" and ext.ok


def test_identified_by_reference_header_when_sheet_renamed():
    ident, ext = run(ds_book(DS_HEADERS, [ds_row("T-1")], sheet="Export"))
    assert (ident.source_type, ident.sheet_name) == ("MDS", "Export") and ext.ok


def test_two_datasheet_like_sheets_are_ambiguous():
    book = make_workbook({"DATASHEET": [DS_HEADERS, ds_row("T-1")], "Copy": [DS_HEADERS, ds_row("T-2")]})
    ident, ext = run(book)
    assert ext is None and codes(ident.issues) == ["SHEET_AMBIGUOUS"]


def test_adapter_is_identification_only():
    assert DATASHEET.source_type == "MDS" and DATASHEET.config.sheet_names == ("DATASHEET",)
