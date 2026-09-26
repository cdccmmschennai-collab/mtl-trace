"""SPIR adapter: identification, shared field discovery, extraction, and the real-file oracle.

Column positions of the real file appear ONLY here, as fixture assertions — never in the adapter.
"""

import io
from collections import Counter

import openpyxl
import pytest

from backend.domain.sources import Attribute
from backend.extraction import registry
from backend.extraction.tabular import load_workbook
from tests.conftest import (
    ORIGINAL_BACKUP_FILE,
    REAL_DATASHEET_FILE,
    REAL_SPIR_FILE,
    make_workbook,
    without_column,
)


def run(data: bytes, name: str = "spir.xlsx", source_type: str | None = None):
    book = load_workbook(name, data)
    ident = registry.identify(book, source_type)
    return ident, (registry.extract(book, ident) if ident.source_type else None)


def by_tag(ext) -> dict[str, dict[Attribute, str | None]]:
    return {r.normalized_key: {a: v.raw for a, v in r.values.items()} for r in ext.rows}


def independent_spir(data: bytes) -> list[dict]:
    """Read SPIR DATA with plain openpyxl and stripped-header lookup — a separate code path."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = list(wb["SPIR DATA"].iter_rows(values_only=True))
    wb.close()
    idx = {str(h).strip(): i for i, h in enumerate(rows[0]) if h is not None}
    out = []
    for r in rows[1:]:
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in r):
            continue
        out.append({k: r[i] for k, i in idx.items()})
    return out


def filled(v) -> bool:
    return v is not None and str(v).strip() != ""


# ---------------------------------------------------------------- real file

def test_real_spir_identified_and_tracker_sheet_skipped(real_spir_bytes):
    ident, ext = run(real_spir_bytes, REAL_SPIR_FILE.name)
    assert (ident.source_type, ident.sheet_name) == ("SPIR", "SPIR DATA")
    assert [s["sheet"] for s in ident.skipped_sheets] == ["PRD-COUNT"]
    assert ext.ok and ext.issues == []
    assert ext.provided_attributes == tuple(Attribute)


def test_each_real_file_is_recognised_as_its_own_type(real_spir_bytes, real_datasheet_bytes):
    """Auto-identification with both adapters registered: no cross-recognition."""
    assert run(real_spir_bytes, REAL_SPIR_FILE.name)[0].source_type == "SPIR"
    assert run(real_datasheet_bytes, REAL_DATASHEET_FILE.name)[0].source_type == "MDS"
    ident, ext = run(real_datasheet_bytes, REAL_DATASHEET_FILE.name, source_type="SPIR")
    assert ext is None and [i.code for i in ident.issues] == ["NOT_THIS_SOURCE_TYPE"]


def test_real_spir_mapping_fixture_positions(real_spir_bytes):
    """Positions verified in the real file. Asserted here; never encoded anywhere in backend/."""
    _, ext = run(real_spir_bytes, REAL_SPIR_FILE.name)
    cols = [(b.key, b.column, b.header) for b in ext.mapping]
    assert cols == [
        ("TAG_KEY", "B", "TAG NUMBER"), ("TAG_AS_PER_DOC", "F", "SPIR TAG NUMBER"), ("MAKE", "N", "MAKE"),
        ("MODEL", "O", "MODEL"), ("SERIAL_NUMBER", "P", "SERIAL NUMBER"), ("PART_NUMBER", "Q", "PART NUMBER"),
        ("DOCUMENT_REFERENCE", "H", "SPIR REF-INITIAL"), ("DOCUMENT_REFERENCE", "J", "SPIR REF-NORMAL"),
        ("DOCUMENT_REFERENCE", "L", "SPIR REF-LCS"),
        ("DOCUMENT_IDB", "G", "DOKNR INITIAL"), ("DOCUMENT_IDB", "I", "DOKNR NORMAL"), ("DOCUMENT_IDB", "K", "DOKNR LCS"),
        ("PAGE_NO", "M", "PAGE NO"),
    ]
    # Everything else in the sheet is reported as found and left alone.
    unbound = set(ext.headers_found) - {b.header for b in ext.mapping}
    assert {"EQUIPMENT DESCRIPTION", "SIZE & RATING", "VENDOR NAME", "COUNTRY", "REMARKS"} <= unbound


def test_real_spir_counts_derived_from_file(real_spir_bytes):
    expected = independent_spir(real_spir_bytes)
    _, ext = run(real_spir_bytes, REAL_SPIR_FILE.name)
    assert len(ext.rows) == len(expected)
    counts = Counter(a for r in ext.rows for a, v in r.values.items() if v.normalized is not None)
    for attr, header in [(Attribute.TAG_NUMBER, "SPIR TAG NUMBER"), (Attribute.MAKE, "MAKE"),
                         (Attribute.MODEL, "MODEL"), (Attribute.SERIAL_NUMBER, "SERIAL NUMBER"),
                         (Attribute.PART_NUMBER, "PART NUMBER")]:
        assert counts.get(attr, 0) == sum(1 for e in expected if filled(e[header])), attr


def test_real_spir_regression_figures(real_spir_bytes):
    _, ext = run(real_spir_bytes, REAL_SPIR_FILE.name)
    counts = Counter(a for r in ext.rows for a, v in r.values.items() if v.normalized is not None)
    assert len(ext.rows) == 464 and len({r.normalized_key for r in ext.rows}) == 464
    assert [counts.get(a, 0) for a in Attribute] == [377, 458, 355, 331, 370]
    tag_vals = [(r.normalized_key, r.values[Attribute.TAG_NUMBER].normalized) for r in ext.rows
                if r.values[Attribute.TAG_NUMBER].normalized]
    assert sum(1 for k, v in tag_vals if v.upper() == k) == 357
    assert sum(1 for k, v in tag_vals if v.upper() != k) == 20


def test_join_key_and_spir_tag_number_are_not_conflated(real_spir_bytes):
    _, ext = run(real_spir_bytes, REAL_SPIR_FILE.name)
    rows = {r.normalized_key: r for r in ext.rows}
    assert rows["68-FIT-1987"].values[Attribute.TAG_NUMBER].normalized == "68-FT-1987"
    assert rows["80-FGP-02S-1"].values[Attribute.TAG_NUMBER].normalized == "80-F&G-02S-1"


def test_every_spir_stage_reference_is_kept(real_spir_bytes):
    expected = independent_spir(real_spir_bytes)
    _, ext = run(real_spir_bytes, REAL_SPIR_FILE.name)
    stages = ("INITIAL", "NORMAL", "LCS")
    for row, e in zip(ext.rows, expected):
        loc = ext.location_of(row)
        refs = [str(e[f"SPIR REF-{s}"]).strip() for s in stages if filled(e[f"SPIR REF-{s}"])]
        idbs = [str(e[f"DOKNR {s}"]).strip() for s in stages if filled(e[f"DOKNR {s}"])]
        assert loc["document_reference"] == ("; ".join(refs) or None)
        assert loc["document_idb"] == ("; ".join(idbs) or None)
        assert loc["page"] == (str(e["PAGE NO"]) if filled(e["PAGE NO"]) else None)
        for s in stages:
            if filled(e[f"SPIR REF-{s}"]):
                assert loc["source_fields"][f"SPIR REF-{s}"] == str(e[f"SPIR REF-{s}"]).strip()
    assert any(loc_refs.count(";") == 2 for loc_refs in (ext.location_of(r)["document_reference"] for r in ext.rows))


def test_raw_values_preserved(real_spir_bytes):
    _, ext = run(real_spir_bytes, REAL_SPIR_FILE.name)
    rows = {r.normalized_key: r for r in ext.rows}
    serial = rows["6834-P-64A"].values[Attribute.SERIAL_NUMBER]
    assert (serial.raw, serial.normalized, serial.raw_type, serial.column) == ("665436", "665436", "int", "P")


# The one cell where ORIGINAL BACKUP does not reproduce the SPIR verbatim: the source reads
# 'JIANGSU YDF VALVE CO.,LTD'; the human backup has 'JIANGSU YDF VALVE CO,LTD'. The source value is kept
# (see docs/SOURCE_ADAPTERS.md). Any other difference fails.
KNOWN_SPIR_BACKUP_DEVIATIONS = {("MAKE COMPARISION", "88-GV-1601"): ("JIANGSU YDF VALVE CO.,LTD", "JIANGSU YDF VALVE CO,LTD")}
SHEETS = {"TAG NUMBER COMPARISION": Attribute.TAG_NUMBER, "MAKE COMPARISION": Attribute.MAKE,
          "MODEL COMPARISION": Attribute.MODEL, "SERIAL NUMBER COMPARISION": Attribute.SERIAL_NUMBER,
          "PART NUMBER COMPARISION": Attribute.PART_NUMBER}


def test_extraction_reproduces_original_backup_spir_columns(real_spir_bytes):
    """Oracle: every SPIR cell of ORIGINAL BACKUP equals the extracted raw value for that tag (or both empty)."""
    if not ORIGINAL_BACKUP_FILE.exists():
        pytest.skip("ORIGINAL BACKUP not present")
    _, ext = run(real_spir_bytes, REAL_SPIR_FILE.name)
    extracted = by_tag(ext)
    wb = openpyxl.load_workbook(ORIGINAL_BACKUP_FILE, read_only=True, data_only=True)
    seen = values = 0
    try:
        for sheet, attr in SHEETS.items():
            rows = list(wb[sheet].iter_rows(values_only=True))
            [col] = [i for i, h in enumerate(rows[0]) if h and "SPIR" in str(h)]
            compared = 0
            for r in rows[1:]:
                if r[1] is None:
                    continue
                oracle = None if r[col] is None or str(r[col]) == "" else str(r[col])
                ours = extracted.get(str(r[1]).strip().upper(), {}).get(attr)
                if (sheet, r[1]) in KNOWN_SPIR_BACKUP_DEVIATIONS:
                    assert (ours, oracle) == KNOWN_SPIR_BACKUP_DEVIATIONS[(sheet, r[1])]
                    seen += 1
                else:
                    assert ours == oracle, (sheet, r[1], ours, oracle)
                compared += 1
                values += oracle is not None
            assert compared == 717
        assert seen == len(KNOWN_SPIR_BACKUP_DEVIATIONS)
        assert values == 377 + 458 + 355 + 331 + 370
    finally:
        wb.close()


def test_reordered_real_columns_extract_identically(real_spir_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(real_spir_bytes), data_only=True)
    ws = wb["SPIR DATA"]
    grid = [[c.value for c in row] for row in ws.iter_rows()]
    wb.remove(ws)
    new = wb.create_sheet("SPIR DATA", 0)
    for row in grid:
        new.append(list(reversed(row)))
    buf = io.BytesIO()
    wb.save(buf)
    _, original = run(real_spir_bytes, REAL_SPIR_FILE.name)
    _, reordered = run(buf.getvalue(), "reordered.xlsx")
    assert reordered.ok and by_tag(reordered) == by_tag(original)
    assert {b.key: b.column for b in original.mapping}["MAKE"] != {b.key: b.column for b in reordered.mapping}["MAKE"]


# ---------------------------------------------------------------- available-data principle

def test_real_spir_without_a_column_still_processes(real_spir_bytes):
    data = without_column(real_spir_bytes, "SPIR DATA", "PART NUMBER")
    _, ext = run(data, "spir-no-part.xlsx")
    assert ext.ok and len(ext.rows) == 464
    assert [(i.code, i.value) for i in ext.issues] == [("FIELD_NOT_PROVIDED", "PART NUMBER")]
    assert ext.provided_attributes == (Attribute.TAG_NUMBER, Attribute.MAKE, Attribute.MODEL, Attribute.SERIAL_NUMBER)
    assert all(Attribute.PART_NUMBER not in r.values for r in ext.rows)
    _, full = run(real_spir_bytes, REAL_SPIR_FILE.name)
    assert {k: {a: v for a, v in vals.items() if a != Attribute.PART_NUMBER} for k, vals in by_tag(full).items()} \
        == by_tag(ext)


def test_sources_with_different_attribute_sets_are_all_valid():
    """TAG+MAKE+MODEL only; and a source with extra, non-MTL columns — both extract what exists."""
    small = make_workbook({"SPIR DATA": [["TAG NUMBER", "SPIR TAG NUMBER", "MAKE", "MODEL"], ["T-1", "T-1", "ACME", "M1"]]})
    _, ext = run(small)
    assert ext.ok and ext.provided_attributes == (Attribute.TAG_NUMBER, Attribute.MAKE, Attribute.MODEL)
    assert by_tag(ext) == {"T-1": {Attribute.TAG_NUMBER: "T-1", Attribute.MAKE: "ACME", Attribute.MODEL: "M1"}}

    wide = make_workbook({"SPIR DATA": [
        ["TAG NUMBER", "MANUFACTURER", "MODEL NO", "SERIAL NO", "PART NO", "SIZE & RATING", "PRESSURE CLASS"],
        ["T-1", "ACME", "M1", "S1", "P1", "4\" 150#", "CL150"]]})
    _, ext = run(wide)
    assert ext.ok and ext.provided_attributes == (Attribute.MAKE, Attribute.MODEL, Attribute.SERIAL_NUMBER,
                                                  Attribute.PART_NUMBER)
    assert by_tag(ext)["T-1"] == {Attribute.MAKE: "ACME", Attribute.MODEL: "M1", Attribute.SERIAL_NUMBER: "S1",
                                  Attribute.PART_NUMBER: "P1"}


def test_ambiguous_attribute_column_is_refused():
    data = make_workbook({"SPIR DATA": [["TAG NUMBER", "MAKE", "SPIR MAKE", "MODEL"], ["T-1", "A", "B", "M"]]})
    _, ext = run(data)
    assert not ext.ok and ext.rows == []
    assert [(i.code, i.value) for i in ext.issues] == [("AMBIGUOUS_FIELD", "MAKE")]
